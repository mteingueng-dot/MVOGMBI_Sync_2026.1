from decimal import Decimal as D

import openpyxl

from mvogmbi_sync import inventaire as inv
from mvogmbi_sync.robot_inventaire import nombre_sage


def clover(i, nom, stock, code="", cout=None):
    return {"id": i, "name": nom, "code": code, "itemStock": {"quantity": stock}, "cost": cout,
            "categories": {"elements": [{"name": "Riz"}]}}


def test_assembler_relie_clover_et_sage():
    sage = {"C1": {"nom": "Riz 5kg", "quantite": D("3.000"), "cout": D("10.0000")},
            "C2": {"nom": "Huile", "quantite": D("0.000"), "cout": None},
            "C3": {"nom": "Sel", "quantite": D("7.000"), "cout": D("1.0000")}}
    lignes = inv.assembler([clover("A", "Riz Parfumé 5kg", 5, "123", 900), clover("B", "Huile 1L", 2),
                            clover("Z", "Sans lien", 1)], sage, {"A": "C1", "B": "C2"})
    assert [(l.id_clover, l.numero_sage, l.quantite_clover, l.quantite_sage) for l in lignes] == [
        ("B", "C2", D(2), D(0)), ("A", "C1", D(5), D(3)), ("Z", "", D(1), None), ("", "C3", None, D(7))]
    assert lignes[1].ecart == D(2) and lignes[1].cout_clover == D(9) and lignes[2].ecart is None


def test_preparer_ramene_les_deux_stocks_a_la_quantite_comptee():
    a = inv.LigneInventaire("A", "", "C1", "Riz", "Riz 5kg", D(5), D(3), D(10), comptee=D(4))
    b = inv.LigneInventaire("B", "", "C1", "Riz bis", "Riz 5kg", D(1), D(3), D(10), comptee=D(1))
    c = inv.LigneInventaire("Z", "", "C2", "Huile", "Huile", D(2), D(0), None, comptee=D(6), cout_clover=D("2.50"))
    d = inv.LigneInventaire("", "", "C3", "", "Sel", None, D(7), D(1), comptee=None)  # non compté : rien
    s = inv.preparer([a, b, c, d])
    assert [(l.id_clover, av, ap) for l, av, ap in s.clover] == [("A", D(5), D(4)), ("Z", D(2), D(6))]  # B inchangé
    assert [(x.numero, x.avant, x.apres, x.cout, x.montant) for x in s.sage] == [
        ("C1", D(3), D(5), D(10), D("20.00")), ("C2", D(0), D(6), D("2.50"), D("15.00"))]
    assert s.problemes == ["C2 : aucun coût dans Sage, coût Clover 2.50 $ utilisé"]
    # stocks relus avant la saisie : l'écart part du stock actuel, pas de celui du fichier
    actuel = [inv.LigneInventaire("A", "", "C1", quantite_clover=D(4), quantite_sage=D(5), cout=D(10))]
    s = inv.preparer([a], actuel)
    assert s.clover == [] and [(x.avant, x.apres) for x in s.sage] == [(D(5), D(4))]


def test_fichier_et_rapport(tmp_path):
    lignes = [inv.LigneInventaire("A", "0123", "C0001", "Riz", "Riz 5kg", D(5), D(3), D("10.5"), "Riz", "31/12/2027")]
    chemin = inv.ecrire(lignes, tmp_path / "inv.xlsx")
    feuille = openpyxl.load_workbook(chemin)[inv.FEUILLE]
    assert [c.value for c in feuille[1]][:6] == ["Code Clover", "Code Sage", "Description Clover", "Description Sage",
                                                 "Quantité Clover", "Quantité Sage"]
    feuille.cell(2, inv.COLONNE_COMPTEE).value = 2
    feuille.parent.save(chemin)
    (l,) = inv.lire_fichier(chemin)
    assert (l.id_clover, l.code_clover, l.numero_sage, l.comptee, l.cout) == ("A", "0123", "C0001", D(2), D("10.5"))
    s = inv.preparer([l])
    rapport = inv.ecrire_rapport(s, tmp_path / "r.xlsx", {"C0001": {"texte": "fait", "montant": "-10.40"}})
    valeurs = list(openpyxl.load_workbook(rapport)["Ajustements Sage"].iter_rows(values_only=True))
    assert valeurs[3][:7] == ("C0001", "Riz 5kg", 3, 2, -1, 10.5, -10.4)  # coût réel lu dans Sage
    assert valeurs[-1][1:7] == ("Coût net de l'ajustement", None, None, None, None, -10.4)


def test_nombres_pour_sage():
    assert nombre_sage(D("4.500"), 3) == "4,5" and nombre_sage(D("-3"), 3) == "-3" and nombre_sage(D("2.0500")) == "2,05"
