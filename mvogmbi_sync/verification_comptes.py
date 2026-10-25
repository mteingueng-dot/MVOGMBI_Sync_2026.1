"""Vérification des comptes associés aux articles de stock dans Sage.

Règle du propriétaire : les comptes d'un article découlent de son compte d'actif (stock) 15xx :
produits (revenus) 40xx et coût des marchandises vendues (CMV) 50xx, avec les mêmes deux derniers chiffres.
Exemple : 1514 Stock - Epices -> 4014 Produit des Ventes - Epices -> 5014 Coût du Stock - Epices.

- Sans compte de produits, Sage bloque l'import de la facture : anomalie « bloquante ».
- Un compte différent de celui attendu n'empêche pas l'import mais classe les ventes au mauvais endroit.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


@dataclass
class ComptesArticle:
    numero: str
    description: str
    actif: str
    produits: str
    cmv: str


@dataclass
class AnomalieCompte:
    article: ComptesArticle
    bloquante: bool
    probleme: str
    proposition: str


def comptes_attendus(actif: str) -> tuple[str, str] | None:
    """(produits, CMV) attendus pour un compte d'actif 15xx ; None si l'actif ne suit pas ce schéma."""
    if len(actif) == 4 and actif.startswith("15"):
        return "40" + actif[2:], "50" + actif[2:]
    return None


def verifier(articles: list[ComptesArticle], noms_comptes: dict[str, str] | None = None) -> list[AnomalieCompte]:
    noms = noms_comptes or {}
    nom = lambda compte: f"{compte} {noms.get(compte, '')}".strip()
    anomalies = []
    for a in articles:
        attendus = comptes_attendus(a.actif)
        if not a.produits:
            proposition = f"Produits : {nom(attendus[0])}" if attendus else "Choisir un compte de produits 40xx"
            anomalies.append(AnomalieCompte(a, True, "Aucun compte de produits (bloque l'import)", proposition))
            continue
        if attendus is None:
            anomalies.append(AnomalieCompte(a, False, f"Compte d'actif {nom(a.actif)} hors stock 15xx",
                                            "Vérifier le compte d'actif de l'article"))
            continue
        produits, cmv = attendus
        if a.produits != produits:
            if comptes_attendus("15" + a.produits[2:]) and a.produits.startswith("40"):
                # Produits cohérents avec une autre catégorie : l'actif est peut-être le compte erroné.
                proposition = (f"Soit produits -> {nom(produits)}, soit actif -> {nom('15' + a.produits[2:])} "
                               f"(et CMV {nom('50' + a.produits[2:])})")
            else:
                proposition = f"Produits -> {nom(produits)}"
            anomalies.append(AnomalieCompte(a, False, f"Produits {nom(a.produits)} ≠ actif {nom(a.actif)}",
                                            proposition))
        elif a.cmv and a.cmv != cmv:
            anomalies.append(AnomalieCompte(a, False, f"CMV {nom(a.cmv)} ≠ actif {nom(a.actif)}",
                                            f"CMV -> {nom(cmv)}"))
    anomalies.sort(key=lambda x: (not x.bloquante, x.article.numero))
    return anomalies


def ecrire_anomalies(anomalies: list[AnomalieCompte], chemin: Path, titre: str = "Comptes à corriger") -> Path:
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = titre[:31]
    feuille.append(["Bloquant", "N° article", "Description", "Actif", "Produits", "CMV", "Problème", "Proposition"])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for x in anomalies:
        a = x.article
        feuille.append(["Oui" if x.bloquante else "", a.numero, a.description, a.actif, a.produits or "(aucun)",
                        a.cmv or "(aucun)", x.probleme, x.proposition])
        if x.bloquante:
            feuille.cell(feuille.max_row, 1).fill = PatternFill("solid", fgColor="FFC7CE")
    for colonne, largeur in zip("ABCDEFGH", (10, 16, 42, 8, 10, 8, 45, 75)):
        feuille.column_dimensions[colonne].width = largeur
    feuille.freeze_panes = "B2"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin
