"""Encaissements des dépôts Clover sur les factures journalières « Vente AAAA-MM-JJ ».

Règles du propriétaire :
- le dépôt Clover dans le compte Desjardins = ventes encaissées moins les espèces, c'est-à-dire le lot
  de cartes que Clover ferme chaque soir (vérifié au cent près : 2026-05-20, 2 740,59 $ - 179,78 $
  d'espèces = lot de 2 560,81 $) ;
- ventes du lundi au jeudi : déposées le lendemain ; ventes du vendredi, samedi et dimanche : déposées
  ensemble le lundi suivant ;
- l'encaissement est enregistré dans Sage à la date du dépôt, au compte Desjardins (1055), et appliqué
  sur la facture de chaque jour de vente pour le montant de son lot ;
- virements Interac reçus et dépôts d'espèces au guichet (relevés Desjardins de la page Rapprochement, règles du
  2026-10-26) : ils font partie des espèces des ventes. Interac : encaissé sur la facture qui couvre le jour du
  virement. Dépôt d'espèces : réparti sur les factures de la période dont il reste des espèces à encaisser (solde
  dû moins les lots de cartes pas encore encaissés), de la plus ancienne à la plus récente. Encaissement daté du
  mouvement de banque (au plus tôt à la date de la facture) ; un seul encaissement par mouvement, pour qu'il se
  rapproche du relevé.

Sage n'accepte pas d'importer un encaissement (Sage_SA_import.exe ne lit que factures, commandes, achats
et écritures de journal) : l'application calcule les encaissements et vérifie dans Sage s'ils sont saisis.
Dans Sage, un encaissement est une ligne tcustrdt de type 1 (montant négatif, compte de banque lBnkAcctId)
rattachée à la facture (lCusTrId).
"""

from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from .periodes import en_millisecondes

CENT = Decimal("0.01")

# États d'un encaissement
APPLIQUE = "Appliqué"
A_SAISIR = "À saisir dans Sage"
ECART = "Montant différent dans Sage"
SANS_FACTURE = "Facture pas encore dans Sage"
A_VENIR = "Dépôt à venir"
A_VERIFIER = "Espèces à répartir à la main"


def date_depot(jour_vente: date) -> date:
    """Lundi à jeudi : le lendemain ; vendredi, samedi, dimanche : le lundi suivant."""
    if jour_vente.weekday() <= 3:
        return jour_vente + timedelta(days=1)
    return jour_vente + timedelta(days=7 - jour_vente.weekday())


def numero_facture(jour_vente: date, prefixe: str = "Vente") -> str:
    return f"{prefixe} {jour_vente.isoformat()}"


def numero_compte_sage(compte: str | int) -> int:
    """« 1055 » -> 10550000 (comptes stockés sur 8 chiffres dans la base Sage)."""
    texte = str(compte).strip()
    return int(texte.ljust(8, "0")) if len(texte) < 8 else int(texte)


@dataclass
class LigneEncaissement:
    jour_vente: date
    montant: Decimal                     # lot(s) de cartes Clover des jours de la facture = part du dépôt
    numero: str = ""                     # n° de facture s'il diffère de « Vente AAAA-MM-JJ » (fin de semaine)
    facture_presente: bool = False
    montant_du: Decimal | None = None    # solde de la facture dans Sage
    deja_encaisse: Decimal = Decimal(0)  # encaissements déjà saisis au compte de dépôt à la date du dépôt
    etat: str = ""
    prefixe: str = "Vente"               # préfixe des factures du point de vente

    @property
    def facture(self) -> str:
        return self.numero or numero_facture(self.jour_vente, self.prefixe)


@dataclass
class Depot:
    date: date
    lignes: list[LigneEncaissement] = field(default_factory=list)
    etat: str = ""
    origine: str = ""      # mouvement du relevé (Interac, dépôt d'espèces) ; « » = lot de cartes Clover
    remarque: str = ""

    @property
    def montant(self) -> Decimal:
        return sum((l.montant for l in self.lignes), Decimal(0))


def lots_par_jour(client, debut: date, fin: date) -> dict[date, Decimal]:
    """Total des lots de cartes Clover fermés chaque jour (ventes - remboursements, hors espèces)."""
    fuseau = ZoneInfo(client.config.fuseau_horaire)
    t0, t1 = en_millisecondes(debut, fin, client.config.fuseau_horaire)
    totaux: dict[date, Decimal] = defaultdict(Decimal)
    for lot in client._pages("batches", {"filter": [f"createdTime>={t0}", f"createdTime<{t1}"]}):
        jour = datetime.fromtimestamp(lot["createdTime"] / 1000, fuseau).date()
        totaux[jour] += Decimal(int(lot.get("totalBatchAmount") or 0)) / 100
    return dict(totaux)


def regrouper(lots: dict[date, Decimal], prefixe: str = "Vente") -> list[Depot]:
    """Dépôts (date de dépôt -> jours de vente) à partir des lots de chaque jour."""
    depots: dict[date, Depot] = {}
    for jour in sorted(lots):
        if lots[jour] == 0:
            continue
        d = date_depot(jour)
        depots.setdefault(d, Depot(d)).lignes.append(LigneEncaissement(jour, lots[jour], prefixe=prefixe))
    return [depots[d] for d in sorted(depots)]


def factures_fin_de_semaine(depots: list[Depot], factures_sage: set[str], prefixe: str = "Vente") -> None:
    """Range les lignes vendredi-samedi-dimanche d'un dépôt du lundi sur la facture de Sage qui couvre chaque jour.

    Règle du propriétaire depuis octobre 2026 : une seule facture « Vente AAAA-MM-JJ/JJ » pour la fin de semaine (une
    ligne pour tout le lot). Les anciennes fins de semaine facturées jour par jour, ou en partie (« Vente 2026-10-24 »
    puis « Vente 2026-10-25/25 » pour le dimanche seul), sont encaissées facture par facture. Jour pas encore facturé :
    sa facture journalière si le reste de la fin de semaine l'est jour par jour, sinon la facture de fin de semaine.
    """
    from .facture import jours_facture, numero_fin_de_semaine

    couverture: dict[date, str] = {}
    for numero in sorted(factures_sage, key=lambda n: -len(jours_facture(n, prefixe))):  # la plus précise l'emporte
        for jour in jours_facture(numero, prefixe):
            couverture[jour] = numero
    for depot in depots:
        groupes: dict[str, list[LigneEncaissement]] = {}
        lignes = []
        for l in depot.lignes:
            if l.jour_vente.weekday() < 4:
                lignes.append(l)
                continue
            vendredi = l.jour_vente - timedelta(days=l.jour_vente.weekday() - 4)
            fin_de_semaine = [vendredi + timedelta(days=i) for i in range(3)]
            numero = couverture.get(l.jour_vente) or (
                numero_facture(l.jour_vente, prefixe) if any(j in couverture for j in fin_de_semaine)  # jour par jour
                else numero_fin_de_semaine(vendredi, fin_de_semaine[-1], prefixe))
            if numero not in groupes:
                groupes[numero] = []
                lignes.append(numero)  # place de la ligne regroupée
            groupes[numero].append(l)
        depot.lignes = [
            l if isinstance(l, LigneEncaissement) else
            (groupes[l][0] if l == numero_facture(groupes[l][0].jour_vente, prefixe) else
             LigneEncaissement(groupes[l][0].jour_vente, sum((x.montant for x in groupes[l]), Decimal(0)), l,
                               prefixe=prefixe))
            for l in lignes]


def noms_factures_sage(base, client_sage: str, prefixe: str = "Vente") -> set[str]:
    return {l["s"] for l in base.requete(
        "SELECT TRIM(t.sSource) s FROM tcustr t JOIN tcustomr c ON c.lId = t.lCusId "
        "WHERE c.sName = %s AND t.sSource LIKE %s", (client_sage, prefixe + " %"))}


def etat_factures_sage(base, client_sage: str, compte_depot: int, factures: list[str]) -> dict[str, dict]:
    """Pour chaque facture : solde dû et encaissements au compte de dépôt, par date."""
    if not factures:
        return {}
    marques = ",".join(["%s"] * len(factures))
    lignes = base.requete(
        "SELECT t.sSource facture, d.nTranType type, d.dtDate date, d.dAmount montant, d.lBnkAcctId banque "
        "FROM tcustr t JOIN tcustomr c ON c.lId = t.lCusId JOIN tcustrdt d ON d.lCusTrId = t.lId "
        f"WHERE c.sName = %s AND t.sSource IN ({marques})", (client_sage, *factures))
    etat: dict[str, dict] = {}
    for l in lignes:
        e = etat.setdefault(l["facture"], {"du": Decimal(0), "encaisse": defaultdict(Decimal),
                                           "recus": defaultdict(list)})
        montant = Decimal(str(l["montant"])).quantize(CENT)
        e["du"] += montant
        if l["type"] == 1 and int(l["banque"] or 0) == compte_depot:
            jour = l["date"].date() if hasattr(l["date"], "date") else l["date"]
            e["encaisse"][jour] -= montant
            e["recus"][jour].append(-montant)
    return etat


def evaluer(depots: list[Depot], etat_sage: dict[str, dict], aujourdhui: date,
            dates_especes: set[date] = frozenset()) -> None:
    """Fixe l'état de chaque ligne et de chaque dépôt (lots de cartes Clover) d'après Sage.

    `dates_especes` : dates où Sage a aussi des encaissements d'Interac ou d'espèces ; un autre montant ce jour-là
    n'est alors pas un écart, et le lot n'est appliqué que si Sage a un encaissement de son montant exact."""
    for depot in depots:
        if depot.origine:
            continue
        for l in depot.lignes:
            e = etat_sage.get(l.facture)
            l.facture_presente = e is not None
            if e is None:
                l.etat = SANS_FACTURE
                continue
            l.montant_du = e["du"]
            l.deja_encaisse = e["encaisse"].get(depot.date, Decimal(0))
            recus = e.get("recus", {}).get(depot.date, [])
            if l.deja_encaisse == l.montant or l.montant in recus:
                l.etat = APPLIQUE
            elif l.deja_encaisse and depot.date not in dates_especes:
                l.etat = ECART
            elif depot.date > aujourdhui:
                l.etat = A_VENIR
            else:
                l.etat = A_SAISIR
        etats = {l.etat for l in depot.lignes}
        for possible in (ECART, SANS_FACTURE, A_SAISIR, A_VENIR, APPLIQUE):
            if possible in etats:
                depot.etat = possible
                break


def calculer(client, base, client_sage: str, compte_depot: int, debut: date, fin: date,
             aujourdhui: date | None = None, racine=None, prefixe: str = "Vente") -> list[Depot]:
    """Dépôts des ventes du `debut` au `fin` inclus, avec leur état dans Sage ; avec `racine`, aussi les virements
    Interac et dépôts d'espèces des relevés Desjardins chargés (page Rapprochement). `client`, `client_sage` et
    `prefixe` : ceux du point de vente ; les espèces du relevé ne vont que sur le point de vente n° 1 (racine=None
    pour les autres)."""
    depots = regrouper(lots_par_jour(client, debut, fin), prefixe)
    if base:
        factures_fin_de_semaine(depots, noms_factures_sage(base, client_sage, prefixe), prefixe)
    banque, dates_especes = mouvements_especes(racine, base, debut, fin) if base and racine else ([], set())
    factures = [l.facture for d in depots for l in d.lignes]
    evaluer(depots, etat_factures_sage(base, client_sage, compte_depot, factures) if base else {},
            aujourdhui or date.today(), dates_especes)
    if banque:
        depots += especes_a_encaisser(base, client_sage, compte_depot, depots, banque, debut, fin)
    if base and racine:
        depots += depots_manuels(racine, base, client_sage, compte_depot, debut, fin)
    depots.sort(key=lambda d: (d.date, bool(d.origine)))
    return depots


class RegistreEncaissementsManuels:
    """Encaissements décidés à la main depuis la page Rapprochement (donnees/encaissements_manuels.json) : un
    mouvement de banque (clé) -> facture de vente et montant ; saisis ensuite dans Sage par la page Encaissements."""

    def __init__(self, racine):
        self.chemin = Path(racine) / "donnees" / "encaissements_manuels.json"
        self.liste: list[dict] = (json.loads(self.chemin.read_text(encoding="utf-8"))
                                  if self.chemin.exists() else [])

    def cles(self) -> set[str]:
        return {e["cle"] for e in self.liste}

    def trouver(self, cle: str) -> dict | None:
        return next((e for e in self.liste if e["cle"] == cle), None)

    def ajouter(self, cle: str, jour: date, montant: Decimal, facture: str, origine_: str) -> None:
        self.liste = [e for e in self.liste if e["cle"] != cle]
        self.liste.append({"cle": cle, "date": jour.isoformat(), "montant": str(montant), "facture": facture,
                           "origine": origine_})

    def retirer(self, cle: str) -> None:
        self.liste = [e for e in self.liste if e["cle"] != cle]

    def enregistrer(self) -> None:
        self.chemin.parent.mkdir(parents=True, exist_ok=True)
        self.chemin.write_text(json.dumps(self.liste, ensure_ascii=False, indent=1), encoding="utf-8")


def depots_manuels(racine, base, client_sage: str, compte_depot: int, debut: date, fin: date) -> list[Depot]:
    """Encaissements manuels : à saisir tant que Sage n'a pas un encaissement de ce montant sur la facture ; les
    encaissements déjà faits ne sont montrés que dans la période."""
    from .rapprochement import factures_ventes, jours_couverts

    liste = RegistreEncaissementsManuels(racine).liste
    if not liste:
        return []
    dates = factures_ventes(base, client_sage)
    etat = etat_factures_sage(base, client_sage, compte_depot, sorted({e["facture"] for e in liste}))
    resultat = []
    for e in liste:
        jour, montant, numero = date.fromisoformat(e["date"]), Decimal(e["montant"]), e["facture"]
        depot = Depot(max(jour, dates.get(numero, jour)), origine=f"Manuel — {e['origine']}")
        if numero not in dates:
            depot.etat, depot.remarque = A_VERIFIER, f"facture {numero} introuvable dans Sage"
            depot.lignes = [LigneEncaissement(jour, montant, numero, etat=A_VERIFIER)]
        else:
            e_sage = etat.get(numero, {})
            fait = montant in e_sage.get("recus", {}).get(depot.date, [])
            if fait and not debut <= jour <= fin + timedelta(days=7):
                continue
            depot.etat = APPLIQUE if fait else A_SAISIR
            depot.lignes = [LigneEncaissement(jours_couverts(numero)[0], montant, numero, True,
                                              e_sage.get("du"), etat=depot.etat)]
        resultat.append(depot)
    return resultat


def mouvements_especes(racine, base, debut: date, fin: date):
    """Virements Interac reçus et dépôts d'espèces du `debut` au `fin` dans les relevés chargés, rapprochés de Sage ;
    renvoie (mouvements, dates des écritures Sage déjà reliées à l'un d'eux)."""
    from .rapprochement import (A_ENCAISSER, RAPPROCHE, RegistreReleves, est_especes, mouvements_sage, periode_des,
                                rapprocher)

    releves = RegistreReleves(racine).lire()
    periode = periode_des(releves)
    if not periode:
        return [], set()
    banque = [m for r in releves for m in r.mouvements]
    sage = mouvements_sage(base, periode[0] - timedelta(days=7), periode[1] + timedelta(days=7))
    rapprocher(banque, sage)
    manuels = RegistreEncaissementsManuels(racine).cles()  # encaissés à la main : pas de répartition automatique
    retenus = [m for m in banque if est_especes(m) and m.etat in (RAPPROCHE, A_ENCAISSER) and m.cle not in manuels
               and debut <= date.fromisoformat(m.date) <= fin]
    dates_sage = {s.ident: date.fromisoformat(s.date) for s in sage}
    return retenus, {dates_sage[n] for m in retenus for n in m.sage if n in dates_sage}


def origine(m) -> str:
    """« Interac ALICE PATRICIA » ou « Dépôt d'espèces »."""
    from .rapprochement import est_interac

    if not est_interac(m):
        return "Dépôt d'espèces"
    morceaux = [x.strip() for x in m.description.split("/") if x.strip()]
    return f"Interac {morceaux[1]}" if len(morceaux) > 1 else "Interac"


def especes_a_encaisser(base, client_sage: str, compte_depot: int, depots_clover: list[Depot], mouvements,
                        debut: date, fin: date) -> list[Depot]:
    """Un encaissement par virement Interac / dépôt d'espèces, réparti sur les factures de vente (voir l'en-tête)."""
    from .rapprochement import factures_ventes, jours_couverts

    factures = {}
    for n, d in factures_ventes(base, client_sage).items():
        jours = jours_couverts(n)
        if jours and jours[0] <= fin and jours[-1] >= debut:
            factures[n] = d
    etat = etat_factures_sage(base, client_sage, compte_depot, list(factures))
    return repartir_especes(mouvements, factures, {n: etat.get(n, {}).get("du", Decimal(0)) for n in factures},
                            depots_clover)


def repartir_especes(mouvements, factures: dict[str, date], dus: dict[str, Decimal],
                     depots_clover: list[Depot]) -> list[Depot]:
    """Répartition (sans Sage) : `factures` n° -> date, `dus` n° -> solde dû dans Sage."""
    from .rapprochement import RAPPROCHE, est_interac, facture_du_jour, jours_couverts

    a_venir: dict[str, Decimal] = defaultdict(Decimal)  # lots de cartes pas encore encaissés, par facture
    for d in depots_clover:
        for l in d.lignes:
            if l.etat != APPLIQUE:
                a_venir[l.facture] += l.montant
    dispo = {n: dus.get(n, Decimal(0)) - a_venir[n] for n in factures}
    resultat = []
    for m in sorted(mouvements, key=lambda x: (x.date, x.rang)):
        jour = date.fromisoformat(m.date)
        depot = Depot(jour, origine=origine(m))
        resultat.append(depot)
        if m.etat == RAPPROCHE:
            depot.etat = APPLIQUE
            depot.lignes = [LigneEncaissement(jour, m.montant, "Sage n° " + ", ".join(map(str, m.sage)), True,
                                              etat=APPLIQUE)]
            continue
        parts: list[tuple[str, Decimal]] = []
        if est_interac(m):
            numero = facture_du_jour(jour, factures)
            if not numero:
                depot.remarque = f"aucune facture de vente du {jour} dans Sage"
            elif dispo[numero] < m.montant:
                depot.remarque = f"{numero} : espèces restant à encaisser {dispo[numero]} $, virement {m.montant} $"
            else:
                parts = [(numero, m.montant)]
        else:
            reste = m.montant
            for numero in sorted((n for n in factures if jours_couverts(n)[0] <= jour and dispo[n] > 0),
                                 key=lambda n: (factures[n], n)):
                part = min(reste, dispo[numero])
                parts.append((numero, part))
                reste -= part
                if not reste:
                    break
            if reste:
                parts = []
                depot.remarque = f"les factures n'ont pas assez d'espèces à encaisser ({reste} $ de trop)"
        if not parts:
            depot.etat = A_VERIFIER
            depot.lignes = [LigneEncaissement(jour, m.montant, "?", etat=A_VERIFIER)]
            continue
        for numero, part in parts:
            dispo[numero] -= part
            depot.lignes.append(LigneEncaissement(jours_couverts(numero)[0], part, numero, True,
                                                  dus.get(numero), etat=A_SAISIR))
        depot.date = max([jour] + [factures[n] for n, _ in parts])  # jamais avant la date de la facture
        depot.etat = A_SAISIR
    return resultat


def ecrire_excel(depots: list[Depot], chemin, compte_depot: str) -> None:
    """Liste des encaissements : une ligne par facture, regroupées par dépôt."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Encaissements"
    feuille.append(["Date du dépôt", "Compte de dépôt", "Facture", "Montant à encaisser", "Montant dû (Sage)", "État"])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1A9B3C")
    for d in depots:
        for l in d.lignes:
            feuille.append([d.date, compte_depot, l.facture, float(l.montant),
                            float(l.montant_du) if l.montant_du is not None else None, l.etat])
        feuille.append([d.date, compte_depot, "Total du dépôt", float(d.montant), None, d.etat])
        for cellule in feuille[feuille.max_row]:
            cellule.font = Font(bold=True)
    for colonne, largeur in zip("ABCDEF", (14, 16, 20, 20, 18, 30)):
        feuille.column_dimensions[colonne].width = largeur
    for ligne in feuille.iter_rows(min_row=2):
        ligne[0].number_format = "yyyy-mm-dd"
        ligne[3].number_format = ligne[4].number_format = "# ##0.00 $"
    classeur.save(chemin)


def encaisser_sur_bureau_cache(racine, config_toml: dict, journal=print, factures: list[str] | None = None,
                               point_vente: str | None = None) -> dict:
    """Exécuté sur le bureau caché : ouvre Sage (fermé au départ), encaisse chaque dépôt « À saisir », referme Sage.

    `factures` (traitement de nuit) : seulement les dépôts dont toutes les lignes vont sur ces factures, c'est-à-dire
    les ventes importées pendant l'exécution (règle du propriétaire : une vente bloquée n'est pas encaissée).
    Renvoie {"encaisses": [...], "a_verifier": [...], "erreur": ""} pour le compte rendu.
    """
    from datetime import timedelta

    from . import coffre
    from .__main__ import charger_config, executable_import, ouvrir_base_sage
    from .api_clover import ClientClover
    from .points_vente import choisir
    from .robot_encaissements import appliquer
    from .sage_bureau import SessionSage, executable_sage

    c = config_toml.get("encaissements", {})
    compte = str(c.get("compte_depot", "1055"))
    mode = str(c.get("mode", "Comptant"))
    # date de départ = premier jour de DÉPÔT à encaisser (aucun dépôt antérieur touché) ; les ventes commencent
    # au premier jour dont le dépôt tombe ce jour-là ou après (lundi : ventes depuis le vendredi)
    premier_depot = date.fromisoformat(str(c.get("date_depart") or config_toml.get("automatique", {}).get("date_depart")))
    debut = premier_depot - timedelta(days=3)
    while date_depot(debut) < premier_depot:
        debut += timedelta(days=1)
    fin = date.today() - timedelta(days=1)
    if c.get("date_fin"):  # réglage d'essai : dernier jour de vente à encaisser
        fin = min(fin, date.fromisoformat(str(c["date_fin"])))
    from .facture import jours_facture

    points = [p for p in choisir(config_toml, point_vente) if p.actif]  # « tous », un code ou « 1,3 »
    if factures is not None:  # ventes de la session : seulement les points de vente et les jours de ces factures
        points = [p for p in points if any(jours_facture(n, p.prefixe) for n in factures)]
        jours = [j for n in factures for p in points for j in jours_facture(n, p.prefixe)]
        if jours:
            debut, fin = max(debut, min(jours)), min(fin, max(jours))
    fichier = config_toml.get("sage_import", {}).get("fichier_entreprise")
    bilan = {"encaisses": [], "a_verifier": [], "erreur": ""}
    exe = executable_sage(executable_import())
    if not exe:
        bilan["erreur"] = "Sage50Accounting.exe introuvable."
        return bilan
    if not points:
        return bilan
    session = SessionSage(exe, journal)
    try:
        utilisateur = str(config_toml.get("sage_import", {}).get("utilisateur", "sysadmin"))
        session.ouvrir(fichier, utilisateur, coffre.lire(coffre.SAGE_IMPORT, ""))
        ouverte = False
        for pdv in points:  # un point de vente après l'autre : ses lots Clover sur les factures de son client
            client = ClientClover(charger_config(racine / "config.toml", pdv))
            with ouvrir_base_sage(config_toml, fichier) as base:
                depots = calculer(client, base, pdv.client_sage, numero_compte_sage(compte), debut, fin,
                                  racine=racine if pdv.principal else None, prefixe=pdv.prefixe)
            depots = [d for d in depots if d.date >= premier_depot]
            if factures is not None:
                importees = set(factures)
                depots = [d for d in depots if d.lignes and all(l.facture in importees for l in d.lignes)]
            a_saisir = [d for d in depots if d.etat == A_SAISIR]
            nom = f"{pdv.libelle} : " if len(points) > 1 else ""
            for d in depots:
                if d.etat == ECART:
                    bilan["a_verifier"].append(f"{nom}Dépôt du {d.date} ({d.montant:,.2f} $) : montant différent "
                                               "déjà saisi dans Sage.")
                elif d.etat == A_VERIFIER:
                    bilan["a_verifier"].append(f"{nom}{d.origine} du {d.date} ({d.montant:,.2f} $) : {d.remarque}.")
            journal(f"{nom}{len(depots)} dépôt(s) " + (f"des {len(factures)} vente(s) importée(s) " if factures else
                                                         f"depuis le {premier_depot} ") +
                    f"(ventes du {debut} au {fin}) ; {len(a_saisir)} à encaisser.")
            if not a_saisir:
                continue
            if not ouverte:
                session.ouvrir_encaissements()
                ouverte = True
            for d, r in zip(a_saisir, appliquer(config_toml, a_saisir, pdv.client_sage, compte, mode, journal)):
                texte = (f"{nom}{d.origine or 'Dépôt'} du {r.date} : {r.montant:,.2f} $ au compte {compte} "
                         f"({', '.join(l.facture for l in d.lignes)})")
                (bilan["encaisses"] if r.traite else bilan["a_verifier"]).append(
                    texte + ("" if r.traite else " — NON encaissé : " + " ".join(r.messages)))
    except Exception as exc:
        bilan["erreur"] = str(exc)
        journal(f"Erreur : {exc}")
    finally:
        session.fermer()
    return bilan
