"""Inventaire du stock dans Clover et dans Sage, en un seul fichier, et saisie d'inventaire dans les deux (demande du
propriétaire du 2026-10-28).

Export : une ligne par article Clover relié à son article Sage (correspondance « Confirmé » ou « Automatique »), puis
les articles Clover sans lien et les articles de stock Sage sans lien. Colonnes : code Clover, code Sage, description
Clover, description Sage, quantité Clover, quantité Sage, écart, coût unitaire Sage (CUMP), et « Quantité comptée » à
remplir.
Saisie : la quantité comptée devient le stock des deux systèmes. Clover = quantité comptée de la ligne ; Sage = somme
des quantités comptées des lignes de son n° (plusieurs articles Clover peuvent partager un article Sage). Le rapport
de saisie donne, par article, les stocks avant et après, l'écart et le coût de l'ajustement (écart Sage × CUMP).
Quantité Sage = tinvbyln.dInStock (stock affiché dans la fiche de l'article, tous emplacements).
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

FEUILLE = "Inventaire"
ENTETES = ["Code Clover", "Code Sage", "Description Clover", "Description Sage", "Quantité Clover", "Quantité Sage",
           "Écart (Clover − Sage)", "Coût unitaire Sage", "Quantité comptée", "Catégorie Clover", "UGS (péremption)",
           "ID Clover"]
COLONNE_COMPTEE = ENTETES.index("Quantité comptée") + 1
ZERO = Decimal(0)


@dataclass
class LigneInventaire:
    id_clover: str = ""
    code_clover: str = ""
    numero_sage: str = ""
    description_clover: str = ""
    description_sage: str = ""
    quantite_clover: Decimal | None = None
    quantite_sage: Decimal | None = None
    cout: Decimal | None = None
    categorie: str = ""
    ugs: str = ""
    comptee: Decimal | None = None
    cout_clover: Decimal | None = None  # coût Clover : sert si Sage n'a aucun coût pour une entrée en stock

    @property
    def ecart(self) -> Decimal | None:
        if self.quantite_clover is None or self.quantite_sage is None:
            return None
        return self.quantite_clover - self.quantite_sage


def _d(valeur) -> Decimal:
    return Decimal(str(valeur or 0)).quantize(Decimal("0.001"))


def lire_sage(base) -> dict[str, dict]:
    """N° Sage -> description, quantité en stock (tous emplacements) et CUMP (valeur / quantité, sinon dernier coût)."""
    lignes = base.requete(
        "SELECT TRIM(i.sPartCode) n, TRIM(i.sName) nom, COALESCE(SUM(y.dInStock), 0) q, COALESCE(SUM(y.dCostStk), 0) v, "
        "MAX(y.dLastCost) d FROM tinvent i LEFT JOIN tinvbyln y ON y.lInventId = i.lId WHERE i.bService = 0 "
        "GROUP BY i.lId, i.sPartCode, i.sName")
    resultat = {}
    for l in lignes:
        quantite, valeur, dernier = _d(l["q"]), Decimal(str(l["v"] or 0)), Decimal(str(l["d"] or 0))
        cump = (valeur / quantite) if quantite > 0 and valeur > 0 else (dernier if dernier > 0 else None)
        resultat[l["n"]] = {"nom": l["nom"], "quantite": quantite,
                            "cout": cump.quantize(Decimal("0.0001"), ROUND_HALF_UP) if cump else None}
    return resultat


def assembler(articles_clover: list[dict], sage: dict[str, dict], liens: dict[str, str]) -> list[LigneInventaire]:
    """`liens` : ID Clover -> n° Sage. Lignes reliées, puis Clover seul, puis Sage seul (ordre alphabétique)."""
    from .articles_sage import reparer_texte

    lignes, relies = [], set()
    for a in articles_clover:
        stock = (a.get("itemStock") or {})
        quantite = stock.get("quantity", stock.get("stockCount"))
        cats = (a.get("categories") or {}).get("elements") or []
        numero = liens.get(a["id"], "")
        s = sage.get(numero)
        if s:
            relies.add(numero)
        lignes.append(LigneInventaire(
            a["id"], a.get("code") or "", numero if s else "", a.get("name", ""),
            reparer_texte(s["nom"]) if s else "", None if quantite is None else _d(quantite),
            s["quantite"] if s else None, s["cout"] if s else None, cats[0]["name"] if cats else "", a.get("sku") or "",
            cout_clover=Decimal(a["cost"]) / 100 if a.get("cost") else None))
    lignes.sort(key=lambda l: (not l.numero_sage, l.description_clover.lower()))
    for numero, s in sorted(sage.items(), key=lambda x: x[1]["nom"].lower()):
        if numero not in relies and numero.strip() not in ("", "."):
            lignes.append(LigneInventaire(numero_sage=numero, description_sage=reparer_texte(s["nom"]),
                                          quantite_sage=s["quantite"], cout=s["cout"]))
    return lignes


def lire(racine: Path, config_toml: dict) -> list[LigneInventaire]:
    """Lit Clover (stock de chaque article) et Sage (l'application ouvre Sage s'il est fermé)."""
    from .__main__ import charger_config, lire_sage as lecture_sage
    from .api_clover import ClientClover
    from .couts_sage import correspondance

    clover = ClientClover(charger_config(Path(racine) / "config.toml")).inventaire()
    sage = lecture_sage(config_toml, lire_sage, config_toml.get("sage_import", {}).get("fichier_entreprise"))
    return assembler(clover, sage, correspondance(racine))


# ----------------------------------------------------------------------------- fichier Excel
def ecrire(lignes: list[LigneInventaire], chemin: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    classeur = openpyxl.Workbook()
    f = classeur.active
    f.title = FEUILLE
    f.append(ENTETES)
    for c in f[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E7D32")
    jaune, rouge = PatternFill("solid", fgColor="FFF4D6"), PatternFill("solid", fgColor="FDE2E4")
    nombre = lambda v: None if v is None else float(v)
    for l in lignes:
        f.append([l.code_clover, l.numero_sage, l.description_clover, l.description_sage, nombre(l.quantite_clover),
                  nombre(l.quantite_sage), nombre(l.ecart), nombre(l.cout), nombre(l.comptee), l.categorie, l.ugs,
                  l.id_clover])
        rang = f.max_row
        f.cell(rang, COLONNE_COMPTEE).fill = jaune
        if l.ecart:
            f.cell(rang, ENTETES.index("Écart (Clover − Sage)") + 1).fill = rouge
        for col in (1, 2, 12):
            f.cell(rang, col).number_format = "@"
    validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                error="Quantité comptée : un nombre (0 ou plus).")
    f.add_data_validation(validation)
    validation.add(f"I2:I{max(f.max_row, 2)}")
    for col, largeur in zip("ABCDEFGHIJKL", (16, 16, 40, 40, 13, 13, 14, 14, 15, 20, 14, 16)):
        f.column_dimensions[col].width = largeur
    f.freeze_panes = "A2"
    f.auto_filter.ref = f.dimensions
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin


def lire_fichier(chemin: Path) -> list[LigneInventaire]:
    """Relit un fichier d'inventaire (avec les quantités comptées saisies)."""
    import openpyxl

    f = openpyxl.load_workbook(chemin, data_only=True)[FEUILLE]
    rangs = f.iter_rows(values_only=True)
    entete = [str(x or "") for x in next(rangs)]
    i = {nom: entete.index(nom) for nom in ENTETES if nom in entete}
    if "Quantité comptée" not in i or "Code Sage" not in i:
        raise ValueError("Ce n'est pas un fichier d'inventaire de l'application (colonnes absentes).")

    def texte(r, nom):
        return str(r[i[nom]]).strip() if nom in i and r[i[nom]] is not None else ""

    def nombre(r, nom):
        v = r[i[nom]] if nom in i else None
        if v is None or str(v).strip() == "":
            return None
        try:
            return Decimal(str(v).replace(",", ".").replace(" ", "")).quantize(Decimal("0.001"))
        except Exception:
            raise ValueError(f"« {nom} » : nombre invalide « {v} »")

    lignes = []
    for r in rangs:
        if not any(r):
            continue
        lignes.append(LigneInventaire(
            texte(r, "ID Clover"), texte(r, "Code Clover"), texte(r, "Code Sage"), texte(r, "Description Clover"),
            texte(r, "Description Sage"), nombre(r, "Quantité Clover"), nombre(r, "Quantité Sage"),
            nombre(r, "Coût unitaire Sage"), texte(r, "Catégorie Clover"), texte(r, "UGS (péremption)"),
            nombre(r, "Quantité comptée")))
    return lignes


# ----------------------------------------------------------------------------- saisie d'inventaire
@dataclass
class AjustementSage:
    numero: str
    description: str
    avant: Decimal
    apres: Decimal
    cout: Decimal | None

    @property
    def ecart(self) -> Decimal:
        return self.apres - self.avant

    @property
    def montant(self) -> Decimal | None:
        return None if self.cout is None else (self.ecart * self.cout).quantize(Decimal("0.01"), ROUND_HALF_UP)


@dataclass
class Saisie:
    clover: list[tuple[LigneInventaire, Decimal, Decimal]] = field(default_factory=list)  # ligne, avant, après
    sage: list[AjustementSage] = field(default_factory=list)
    problemes: list[str] = field(default_factory=list)


def preparer(lignes: list[LigneInventaire], actuel: list[LigneInventaire] | None = None) -> Saisie:
    """Ajustements à faire pour que les deux stocks égalent les quantités comptées.

    `actuel` : inventaire relu juste avant la saisie (les stocks ont pu bouger depuis l'export) ; sans lui, les
    quantités du fichier servent de stock « avant ».
    """
    s = Saisie()
    clover_actuel = {l.id_clover: l.quantite_clover for l in actuel or [] if l.id_clover}
    sage_actuel = {l.numero_sage: (l.quantite_sage, l.cout, l.description_sage) for l in actuel or [] if l.numero_sage}
    couts_clover = {l.numero_sage: l.cout_clover for l in list(actuel or []) + lignes if l.numero_sage and l.cout_clover}
    par_sage: dict[str, list[LigneInventaire]] = defaultdict(list)
    for l in lignes:
        if l.comptee is None:
            continue
        if l.comptee < 0:
            s.problemes.append(f"{l.description_clover or l.description_sage} : quantité comptée négative")
            continue
        if l.id_clover:
            avant = clover_actuel.get(l.id_clover, l.quantite_clover) if actuel is not None else l.quantite_clover
            if avant != l.comptee:
                s.clover.append((l, avant if avant is not None else ZERO, l.comptee))
        if l.numero_sage:
            par_sage[l.numero_sage].append(l)
    for numero, groupe in par_sage.items():
        quantite, cout, description = sage_actuel.get(numero, (groupe[0].quantite_sage, groupe[0].cout,
                                                               groupe[0].description_sage))
        cible = sum((l.comptee for l in groupe), ZERO)
        avant = quantite if quantite is not None else ZERO
        if cible != avant:
            if cout is None and couts_clover.get(numero):
                cout = couts_clover[numero]
                s.problemes.append(f"{numero} : aucun coût dans Sage, coût Clover {cout} $ utilisé")
            s.sage.append(AjustementSage(numero, description, avant, cible, cout))
            if cout is None:
                s.problemes.append(f"{numero} : coût unitaire inconnu (ajustement sans montant)")
    return s


def ecrire_rapport(saisie: Saisie, chemin: Path, resultats_sage: dict[str, dict] | None = None,
                   resultats_clover: dict[str, str] | None = None) -> Path:
    """Rapport de saisie d'inventaire : stocks avant et après, écarts et coût de l'ajustement (variation de la
    valeur du stock lue dans Sage après la saisie ; à défaut, estimation écart × coût unitaire)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    classeur = openpyxl.Workbook()
    f = classeur.active
    f.title = "Ajustements Sage"
    f.append([f"Saisie d'inventaire du {datetime.now():%Y-%m-%d %H:%M}"])
    f["A1"].font = Font(bold=True, size=13)
    f.append([])
    entete = ["Code Sage", "Description Sage", "Stock avant", "Stock compté", "Écart", "Coût unitaire (CUMP)",
              "Coût de l'ajustement", "Résultat"]
    f.append(entete)
    for c in f[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E7D32")
    total_plus, total_moins = ZERO, ZERO
    for a in saisie.sage:
        r = (resultats_sage or {}).get(a.numero) or {}
        montant = Decimal(r["montant"]) if r.get("montant") else a.montant
        f.append([a.numero, a.description, float(a.avant), float(a.apres), float(a.ecart),
                  None if a.cout is None else float(a.cout), None if montant is None else float(montant),
                  r.get("texte", "")])
        f.cell(f.max_row, 1).number_format = "@"
        if montant:
            total_plus += max(montant, ZERO)
            total_moins += min(montant, ZERO)
    f.append([])
    for libelle, valeur in (("Ajustements en plus", total_plus), ("Ajustements en moins (pertes)", total_moins),
                            ("Coût net de l'ajustement", total_plus + total_moins)):
        f.append(["", libelle, None, None, None, None, float(valeur)])
        f.cell(f.max_row, 2).font = Font(bold=True)
        f.cell(f.max_row, 7).font = Font(bold=True)
    for col, largeur in zip("ABCDEFGH", (16, 44, 12, 12, 10, 18, 18, 40)):
        f.column_dimensions[col].width = largeur

    g = classeur.create_sheet("Stock Clover")
    g.append(["Code Clover", "Description Clover", "Stock avant", "Stock compté", "Écart", "Résultat", "ID Clover"])
    for c in g[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E7D32")
    for l, avant, apres in saisie.clover:
        g.append([l.code_clover, l.description_clover, float(avant), float(apres), float(apres - avant),
                  (resultats_clover or {}).get(l.id_clover, ""), l.id_clover])
    for col, largeur in zip("ABCDEFG", (16, 44, 12, 12, 10, 40, 16)):
        g.column_dimensions[col].width = largeur
    if saisie.problemes:
        h = classeur.create_sheet("À vérifier")
        for p in saisie.problemes:
            h.append([p])
        h.column_dimensions["A"].width = 90
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin


def nom_fichier(prefixe: str) -> str:
    return re.sub(r"[^\w-]", "_", f"{prefixe}_{datetime.now():%Y%m%d_%H%M}") + ".xlsx"


# ----------------------------------------------------------------------------- saisie complète
def ajuster_sage(racine: Path, config_toml: dict, ajustements: list[AjustementSage], jour: str, source: str = "",
                 remarque: str = "") -> dict[str, dict]:
    """Ajustements Sage par le Sage ouvert sur le fichier configuré, sinon l'application ouvre Sage elle-même.
    Renvoie {n° Sage : {texte, montant}}."""
    import json
    import tempfile

    from .__main__ import commande_moteur
    from .robot_articles import fenetre_principale_ouverte
    from .sage_bureau import executer_sur_bureau_cache, sage_ouvert

    demandes = [{"numero": a.numero, "cible": str(a.apres), "cout": None if a.cout is None else str(a.cout)}
                for a in ajustements]
    if not demandes:
        return {}
    fichier = config_toml.get("sage_import", {}).get("fichier_entreprise")
    if fenetre_principale_ouverte(fichier):
        return {l.numero: resume(l) for l in ajuster_dans_sage(config_toml, demandes, jour, source=source,
                                                               remarque=remarque)}
    if sage_ouvert():
        raise RuntimeError("Sage est ouvert sur un autre fichier d'entreprise : fermez-le, l'application l'ouvrira "
                           "elle-même.")
    with tempfile.TemporaryDirectory() as dossier:
        entree, bilan = Path(dossier) / "ajustements.json", Path(dossier) / "bilan.json"
        entree.write_text(json.dumps({"jour": jour, "lignes": demandes, "source": source, "remarque": remarque}),
                          encoding="utf-8")
        executer_sur_bureau_cache(commande_moteur() + ["inventaire-bureau", "--ajustements", str(entree),
                                                      "--resultat", str(bilan)], racine)
        resultat = json.loads(bilan.read_text(encoding="utf-8")) if bilan.exists() else {}
    if resultat.get("erreur"):
        raise RuntimeError(resultat["erreur"])
    return resultat.get("resultats", {})


def ajuster_dans_sage(config_toml: dict, demandes: list[dict], jour: str, journal=print, source: str = "",
                      remarque: str = ""):
    from .robot_inventaire import LigneAjustement, appliquer

    lignes = [LigneAjustement(d["numero"], Decimal(d["cible"]), Decimal(d["cout"]) if d.get("cout") else None)
              for d in demandes]
    return appliquer(config_toml, lignes, jour, journal, source, remarque)


def resume(l) -> dict:
    """Résultat d'une ligne : texte et coût réel de l'ajustement (variation de la valeur du stock dans Sage)."""
    texte = f"fait : {l.avant} -> {l.apres}" if l.fait else "NON fait : " + " ; ".join(l.messages)
    montant = None
    if l.fait and l.valeur_avant is not None and l.valeur_apres is not None and l.avant != l.apres:
        montant = str(l.valeur_apres - l.valeur_avant)
    return {"texte": texte, "montant": montant}


def saisir(racine: Path, config_toml: dict, chemin: Path, journal=print) -> tuple[Saisie, Path]:
    """Saisie d'inventaire d'un fichier exporté puis rempli : Clover, puis Sage, puis rapport (sortie/inventaire)."""
    from datetime import date

    from .__main__ import charger_config
    from .api_clover import ClientClover

    lignes = lire_fichier(chemin)
    if not any(l.comptee is not None for l in lignes):
        raise ValueError("Aucune quantité comptée dans le fichier (colonne « Quantité comptée »).")
    journal("Lecture des stocks actuels de Clover et de Sage…")
    saisie = preparer(lignes, lire(racine, config_toml))
    resultats_clover = {}
    if saisie.clover:
        client = ClientClover(charger_config(Path(racine) / "config.toml"))
        for l, _avant, apres in saisie.clover:
            try:
                client.regler_stock(l.id_clover, apres)
                resultats_clover[l.id_clover] = "fait"
            except Exception as exc:
                resultats_clover[l.id_clover] = f"NON fait : {exc}"
        journal(f"Clover : {sum(1 for r in resultats_clover.values() if r == 'fait')} stock(s) mis à jour.")
    resultats_sage = {}
    if saisie.sage:
        journal(f"Sage : {len(saisie.sage)} ajustement(s)…")
        try:
            resultats_sage = ajuster_sage(racine, config_toml, saisie.sage, date.today().isoformat())
        except Exception as exc:
            resultats_sage = {a.numero: {"texte": f"NON fait : {exc}", "montant": None} for a in saisie.sage}
    rapport = ecrire_rapport(saisie, Path(racine) / "sortie" / "inventaire" / nom_fichier("Saisie_inventaire"),
                             resultats_sage, resultats_clover)
    return saisie, rapport
