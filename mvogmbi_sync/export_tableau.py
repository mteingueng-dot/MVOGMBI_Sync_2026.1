"""Export de n'importe quel tableau de l'interface (ttk.Treeview) en Excel (.xlsx), en CSV ou en PDF.

Ce qui est exporté = ce qui est affiché (filtre et recherche compris), colonnes dans l'ordre de l'écran ; les
lignes enfants (dépôt -> factures de la page Encaissements) sont décalées sous leur parent. Excel : montants
convertis en nombres (« 1 234,56 $ », « 1,234.56 », « -5.95 ») ; codes, dates et numéros restent du texte.
CSV : séparateur « ; » et UTF-8 avec BOM, comme l'attend Excel en français.
"""

from __future__ import annotations

import csv
import os
import re
import tkinter as tk
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from .langues import t

MONTANT = re.compile(r"^-?\d{1,3}(?:[   ,]\d{3})*(?:[.,]\d+)?$|^-?\d+[.,]\d+$")


def contenu(tableau: ttk.Treeview) -> tuple[list[str], list[list[str]]]:
    """(en-têtes, lignes) du tableau tel qu'il est affiché."""
    colonnes = list(tableau["displaycolumns"])
    if colonnes in (["#all"], [], ("#all",)):
        colonnes = list(tableau["columns"])
    arbre = "tree" in str(tableau["show"])
    entetes = ([tableau.heading("#0", "text")] if arbre else []) + [tableau.heading(c, "text") for c in colonnes]
    lignes: list[list[str]] = []

    def parcourir(parent: str, niveau: int) -> None:
        for ligne in tableau.get_children(parent):
            valeurs = [str(tableau.set(ligne, c)) for c in colonnes]
            if arbre:
                valeurs.insert(0, "    " * niveau + str(tableau.item(ligne, "text")))
            lignes.append(valeurs)
            parcourir(ligne, niveau + 1)

    parcourir("", 0)
    return entetes, lignes


def nombre(texte: str) -> Decimal | None:
    """« 1 234,56 $ » -> 1234.56 ; None si ce n'est pas un montant (code, date, n° de facture)."""
    s = texte.replace("$", "").strip()
    if not MONTANT.match(s):
        return None
    if "," in s and "." in s:
        s = s.replace(",", "")           # 1,234.56
    elif "," in s:
        entier, _, decimales = s.rpartition(",")
        s = entier.replace(",", "") + "." + decimales if len(decimales) != 3 else s.replace(",", "")
    s = re.sub(r"[   ]", "", s)
    if "." not in s and texte.strip().lstrip("-")[:1] == "0" and len(s.lstrip("-")) > 1:
        return None                      # code-barres « 0123456 » : du texte
    if "." not in s and "$" not in texte:
        return None                      # entier sans $ : numéro ou quantité entière laissés tels quels
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def ecrire_csv(chemin: Path, entetes: list[str], lignes: list[list[str]]) -> None:
    with open(chemin, "w", newline="", encoding="utf-8-sig") as f:
        sortie = csv.writer(f, delimiter=";")
        sortie.writerow(entetes)
        sortie.writerows([v.strip() for v in l] for l in lignes)


def ecrire_excel(chemin: Path, titre: str, entetes: list[str], lignes: list[list[str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = re.sub(r"[\\/*?:\[\]]", "-", titre)[:31] or "Tableau"
    feuille.append(entetes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F7A3A")
        cellule.alignment = Alignment(vertical="center")
    for valeurs in lignes:
        rangee = []
        for v in valeurs:
            n = nombre(v)
            rangee.append(float(n) if n is not None else v.rstrip())
        feuille.append(rangee)
        for cellule, v in zip(feuille[feuille.max_row], valeurs):
            n = nombre(v)
            if n is not None:
                decimales = -n.as_tuple().exponent if n.as_tuple().exponent < 0 else 0
                cellule.number_format = "#,##0" + ("." + "0" * decimales if decimales else "")
    for i, entete in enumerate(entetes, 1):
        largeur = max([len(str(entete))] + [len(l[i - 1]) for l in lignes if i - 1 < len(l)])
        feuille.column_dimensions[get_column_letter(i)].width = min(60, largeur + 2)
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions
    classeur.save(chemin)


EDGE = [Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe")]


def ecrire_pdf(chemin: Path, titre: str, entetes: list[str], lignes: list[list[str]], images=None,
               sous_titre: str = "") -> None:
    """PDF du tableau (page HTML imprimée par Microsoft Edge sans fenêtre, comme le manuel : aucune bibliothèque
    de plus) : en-tête du marché, titre, date, en-têtes verts, une ligne sur deux teintée, montants à droite,
    page à l'horizontale au-delà de 7 colonnes (ou s'il y a des graphiques). `images` : images PIL (graphiques,
    cartes d'indicateurs) placées avant le tableau ; sans en-têtes, le PDF ne contient que les images."""
    pdf_depuis_html(page_html(titre, entetes, lignes, images, sous_titre), chemin)


class _ProgrammeExterne:
    """Le programme installé (PyInstaller) donne à ce qu'il lance son propre dossier de bibliothèques
    (SetDllDirectory) : Edge y prenait de mauvaises DLL et s'arrêtait sans rien dire (« Le PDF n'a pas pu être
    créé », constat du 2026-10-29 : aucun PDF de bon de commande créé par le programme installé). Le réglage normal
    est remis le temps de lancer Edge, puis rétabli."""

    def __enter__(self):
        import ctypes
        import sys

        self.avant = None
        if getattr(sys, "frozen", False) and sys.platform == "win32":
            noyau = ctypes.windll.kernel32
            tampon = ctypes.create_unicode_buffer(1024)
            if noyau.GetDllDirectoryW(1024, tampon):
                self.avant = tampon.value
            noyau.SetDllDirectoryW(None)
        return self

    def __exit__(self, *_):
        import ctypes

        if self.avant:
            ctypes.windll.kernel32.SetDllDirectoryW(self.avant)


def pdf_depuis_html(page: str, chemin: Path, delai: int = 120) -> None:
    """Page HTML imprimée en PDF par Microsoft Edge sans fenêtre (tableaux, bons de commande, manuel du profil)."""
    import shutil
    import subprocess
    import tempfile

    edge = next((e for e in EDGE if e.exists()), None)
    if not edge:
        raise RuntimeError(t("Microsoft Edge introuvable : l'export PDF en a besoin."))
    raison = ""
    # des processus d'Edge gardent quelques secondes leur profil ouvert : le dossier temporaire est effacé sans
    # erreur (« WinError 32 … Extension Rules », constat du 2026-10-29 : le PDF était fait, l'erreur venait du
    # ménage) ; ce qui reste est effacé par Windows plus tard
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as dossier:
        # Edge écrit dans un dossier temporaire (chemin simple), puis le PDF est copié à sa place : un dossier
        # OneDrive, réseau ou protégé ne fait plus échouer Edge (constat du 2026-10-28)
        source, pdf = Path(dossier) / "tableau.html", Path(dossier) / "tableau.pdf"
        source.write_text(page, encoding="utf-8")
        for options in (["--headless"], ["--headless=old"]):  # 2e essai avec l'ancien mode sans fenêtre
            try:
                with _ProgrammeExterne():
                    r = subprocess.run([str(edge), *options, "--disable-gpu", "--no-pdf-header-footer",
                                        f"--user-data-dir={Path(dossier) / 'profil'}", f"--print-to-pdf={pdf}",
                                        source.as_uri()], timeout=delai, capture_output=True,
                                       stdin=subprocess.DEVNULL, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                lignes = (r.stderr or b"").decode("utf-8", "replace").strip().splitlines()
                raison = (lignes[-1] if lignes else f"Edge s'est arrêté (code {r.returncode})")[:200]
            except subprocess.TimeoutExpired:
                raison = "Edge n'a pas répondu en 2 minutes"
            if pdf.exists() and pdf.stat().st_size > 0:
                break
        if not pdf.exists():
            raise RuntimeError(t("Le PDF n'a pas pu être créé.") + (f" ({raison})" if raison else ""))
        Path(chemin).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(pdf, chemin)


def page_html(titre: str, entetes: list[str], lignes: list[list[str]], images=None, sous_titre: str = "",
              imprimer: bool = False) -> str:
    """Page HTML du tableau (PDF et impression) ; `imprimer` : la boîte d'impression s'ouvre dès l'affichage.
    En-tête : logo et informations de l'entreprise lues dans Clover (nom, adresse, téléphone, site web)."""
    import html
    from datetime import datetime

    e = html.escape
    entete = entete_html()
    import base64
    import io

    blocs_images = []
    for image in images or []:
        tampon = io.BytesIO()
        image.save(tampon, format="PNG")
        blocs_images.append("<img class='graphique' src='data:image/png;base64,"
                            + base64.b64encode(tampon.getvalue()).decode("ascii") + "'>")
    entier = re.compile(r"^-?\d[\d\s ]*([.,]\d+)?( %)?$")  # nombres, quantités, pourcentages : à droite
    nombres = [any(nombre(l[i]) is not None or entier.match(str(l[i]).strip()) for l in lignes[:50] if i < len(l))
               for i in range(len(entetes))]
    corps = "".join(
        "<tr>" + "".join(f"<td class='{'n' if nombres[i] else ''}'>{e(str(v).strip())}</td>"
                         for i, v in enumerate(l)) + "</tr>" for l in lignes)
    page = f"""<!doctype html><html lang='fr'><head><meta charset='utf-8'><title>{e(titre)}</title><style>
@page {{ size: Letter {'landscape' if len(entetes) > 7 or images else 'portrait'}; margin: 12mm 10mm; }}
img.graphique {{ display: block; max-width: 100%; max-height: 165mm; margin: 6px auto 10px; border: 1px solid #e2e8e4;
                 page-break-inside: avoid; }}
* {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
body {{ font-family: 'Segoe UI', Arial, sans-serif; font-size: 8.5pt; color: #1f2d24; margin: 0; }}
/* en-tête moderne : filet aux couleurs du marché, bandeau clair arrondi, titre au centre */
.filet {{ height: 4px; border-radius: 2px; margin-bottom: 8px;
          background: linear-gradient(90deg, #1a9b3c 0 34%, #d62e3c 34% 67%, #f5c518 67% 100%); }}
.bandeau {{ display: flex; justify-content: space-between; align-items: center; gap: 14px;
            background: #f4f9f5; border: 1px solid #e1ece4; border-radius: 10px; padding: 10px 14px; }}
.ident {{ display: flex; gap: 10px; align-items: center; min-width: 0; flex: 0 1 38%; }}
.ident img {{ height: 52px; width: 52px; object-fit: contain; background: #fff; border-radius: 8px;
              border: 1px solid #e1ece4; padding: 3px; }}
.marche {{ color: #13772d; font-weight: 700; font-size: 11.5pt; letter-spacing: .2px; }}
.coord {{ color: #5b6a60; font-size: 7.5pt; line-height: 1.4; }}
.titre {{ flex: 1; text-align: center; }}
.titre h1 {{ margin: 0; font-size: 16pt; font-weight: 700; color: #1f2d24; letter-spacing: .3px; }}
.titre .sous {{ margin: 3px 0 0; color: #6b7a70; font-size: 8pt; }}
.date {{ flex: 0 0 auto; color: #6b7a70; font-size: 7.5pt; text-align: right; line-height: 1.5; }}
.date b {{ display: block; color: #13772d; font-size: 9pt; }}
.sous {{ color: #6b7a70; }}
table {{ width: 100%; border-collapse: separate; border-spacing: 0; }}
thead {{ display: table-header-group; }}
th {{ background: #1f7a3a; color: #fff; text-align: left; padding: 6px 7px; font-weight: 600; font-size: 8pt; }}
th.n {{ text-align: right; }}
th:first-child {{ border-top-left-radius: 6px; }}
th:last-child {{ border-top-right-radius: 6px; }}
td {{ padding: 4px 7px; border-bottom: 1px solid #e8efe9; vertical-align: top; }}
td.n {{ text-align: right; white-space: nowrap; }}
tr:nth-child(even) td {{ background: #f4f8f5; }}
tr {{ page-break-inside: avoid; }}
</style></head><body>
<div class='filet'></div>
<div class='bandeau'>{entete}
<div class='titre'><h1>{e(titre)}</h1>
<div class='sous'>{e(sous_titre or (t("{n} ligne(s)", n=len(lignes)) if entetes else ""))}</div></div>
<div class='date'><b>{datetime.now():%Y-%m-%d}</b>{datetime.now():%H:%M}<br>MVOGMBI_Sync</div></div>
<div style='height:10px'></div>
{"".join(blocs_images)}
{f'<table><thead><tr>{"".join(f"<th class={chr(39)}{chr(110) if nombres[i] else chr(32)}{chr(39)}>{e(str(x))}</th>" for i, x in enumerate(entetes))}</tr></thead><tbody>{corps}</tbody></table>'
 if entetes else ""}
{"<script>window.addEventListener('load', () => window.print());</script>" if imprimer else ""}
</body></html>"""
    return page


def entete_html() -> str:
    """Bloc gauche de l'en-tête : logo du marché, nom, adresse, téléphone et site web (Clover, gardés localement)."""
    import base64
    import html

    from .design import LOGO
    from .entreprise import Entreprise, lire

    e = html.escape
    info = lire(RACINE_ENTETE) if RACINE_ENTETE else Entreprise()
    logo = ""
    try:
        logo = ("<img src='data:image/png;base64," + base64.b64encode(Path(LOGO).read_bytes()).decode("ascii")
                + "'>")
    except OSError:
        pass
    coordonnees = [x for x in (info.adresse, info.ville,
                               "  •  ".join(y for y in (f"Tél. {info.telephone}" if info.telephone else "", info.site)
                                            if y)) if x]
    return (f"<div class='ident'>{logo}<div><div class='marche'>{e(info.nom)}</div>"
            f"<div class='coord'>{'<br>'.join(e(x) for x in coordonnees)}</div></div></div>")


def imprimer(tableau: ttk.Treeview, nom) -> None:
    """Imprime le tableau affiché (toutes ses lignes) : la page s'ouvre dans le navigateur avec la boîte
    d'impression de Windows (choix de l'imprimante, aperçu, orientation)."""
    import tempfile
    from datetime import datetime

    if not autorise(tableau):
        return
    nom = _nom(nom)
    entetes, lignes = contenu(tableau)
    if not lignes:
        messagebox.showinfo(nom, t("Le tableau est vide : rien à imprimer."), parent=tableau)
        return
    dossier = Path(tempfile.gettempdir()) / "MVOGMBI_Sync_impression"
    dossier.mkdir(exist_ok=True)
    for ancien in dossier.glob("*.html"):  # impressions précédentes : effacées
        try:
            ancien.unlink()
        except OSError:
            pass
    chemin = dossier / f"impression_{datetime.now():%Y%m%d_%H%M%S}.html"
    chemin.write_text(page_html(nom, entetes, lignes, imprimer=True), encoding="utf-8")
    os.startfile(chemin)


def exporter(tableau: ttk.Treeview, nom, format_: str = "xlsx") -> None:
    """Demande où enregistrer puis écrit le fichier ; propose de l'ouvrir."""
    if not autorise(tableau):
        return
    nom = _nom(nom)
    entetes, lignes = contenu(tableau)
    if not lignes:
        messagebox.showinfo(nom, t("Le tableau est vide : rien à exporter."), parent=tableau)
        return
    types = {"xlsx": ("Excel", "*.xlsx"), "csv": ("CSV", "*.csv"), "pdf": ("PDF", "*.pdf")}
    chemin = filedialog.asksaveasfilename(parent=tableau, defaultextension="." + format_,
                                          initialfile=f"{nom_de_fichier(nom)} {date.today().isoformat()}.{format_}",
                                          filetypes=[types[format_]] + [v for k, v in types.items() if k != format_])
    if not chemin:
        return
    chemin = Path(chemin)
    try:
        if chemin.suffix.lower() == ".csv":
            ecrire_csv(chemin, entetes, lignes)
        elif chemin.suffix.lower() == ".pdf":
            tableau.configure(cursor="watch")
            tableau.update_idletasks()
            try:
                ecrire_pdf(chemin, nom, entetes, lignes)
            finally:
                tableau.configure(cursor="")
        else:
            ecrire_excel(chemin.with_suffix(".xlsx"), nom, entetes, lignes)
            chemin = chemin.with_suffix(".xlsx")
    except PermissionError:
        messagebox.showerror(nom, t("Fichier ouvert dans Excel : fermez-le puis recommencez."), parent=tableau)
        return
    except Exception as exc:  # PDF : Edge absent ou en échec
        messagebox.showerror(nom, str(exc), parent=tableau)
        return
    if messagebox.askyesno(nom, t("{n} ligne(s) exportée(s) dans {fichier}.\n\nOuvrir le fichier ?",
                                  n=len(lignes), fichier=chemin.name), parent=tableau):
        os.startfile(chemin)


RACINE_ENTETE: Path | None = None  # dossier de l'application : informations d'entreprise de l'en-tête (Clover)
EXPORT_PERMIS = True  # droit du profil de la personne connectée (réglé par l'application à la connexion)


def export_permis() -> bool:
    return EXPORT_PERMIS


def autorise(parent=None) -> bool:
    """Export permis au profil ; sinon un message (le profil n'a pas le droit d'exporter des données)."""
    if EXPORT_PERMIS:
        return True
    from tkinter import messagebox

    messagebox.showinfo("MVOGMBI_Sync", t("Votre profil n'a pas le droit d'exporter des données."), parent=parent)
    return False


def nom_de_fichier(nom: str) -> str:
    """Titre sans les caractères interdits dans un nom de fichier Windows (« : », « / », « ? »…)."""
    return re.sub(r'[<>:"/\\|?*]+', "-", nom).strip(" -.")[:80] or "MVOGMBI_Sync"


def _nom(nom) -> str:
    """Titre du document : un texte, ou une fonction appelée au moment de l'export (titre des résultats affichés)."""
    texte = nom() if callable(nom) else nom
    return " ".join(str(texte or "").split()) or "MVOGMBI_Sync"


def menu(tableau: ttk.Treeview, nom) -> tk.Menu:
    m = tk.Menu(tableau, tearoff=0)
    m.add_command(label="📗  " + t("Exporter en Excel (.xlsx)…"), command=lambda: exporter(tableau, nom, "xlsx"))
    m.add_command(label="📄  " + t("Exporter en CSV (.csv)…"), command=lambda: exporter(tableau, nom, "csv"))
    m.add_command(label="📕  " + t("Exporter en PDF (.pdf)…"), command=lambda: exporter(tableau, nom, "pdf"))
    m.add_separator()
    m.add_command(label="🖨  " + t("Imprimer…"), command=lambda: imprimer(tableau, nom))
    return m


def rendre_exportable(tableau: ttk.Treeview, nom: str) -> None:
    """Clic droit sur le tableau : Exporter en Excel / en CSV / en PDF."""
    tableau._nom_export = nom
    if not EXPORT_PERMIS:  # profil sans le droit d'exporter : pas de menu d'export
        return
    tableau.bind("<Button-3>", lambda e: menu(tableau, nom).tk_popup(e.x_root, e.y_root), add="+")


def bouton_exporter(parent, tableau: ttk.Treeview, nom: str, **options):
    """Liste déroulante « Exporter ▾ » : Excel, CSV, PDF et Imprimer ; aussi par clic droit sur le tableau."""
    from .design import Bouton, fond_de

    rendre_exportable(tableau, nom)
    if not EXPORT_PERMIS:  # profil sans le droit d'exporter : aucun bouton (cadre vide, sans taille)
        return tk.Frame(parent, width=1, height=1, background=fond_de(parent))
    bouton = None

    def ouvrir() -> None:
        menu(tableau, nom).tk_popup(bouton.winfo_rootx(), bouton.winfo_rooty() + bouton.winfo_height())

    bouton = Bouton(parent, "⬇  " + t("Exporter") + "  ▾", ouvrir, "secondaire", **options)
    return bouton


def placer_a_droite(widget, **options):
    """Place le bouton « Exporter ▾ » le plus à droite de sa barre (demande du propriétaire du 2026-10-28), même si
    d'autres boutons y sont déjà rangés à droite."""
    a_droite = [w for w in widget.master.pack_slaves() if w is not widget and w.pack_info().get("side") == "right"]
    widget.pack(side="right", **({"before": a_droite[0]} if a_droite else {}), **options)
    return widget


def image_widget(widget):
    """Image (PIL) de ce que montre un élément de la fenêtre : graphique, cartes d'indicateurs…"""
    from PIL import ImageGrab

    widget.update()
    x, y = widget.winfo_rootx(), widget.winfo_rooty()
    return ImageGrab.grab((x, y, x + widget.winfo_width(), y + widget.winfo_height()), all_screens=True)


def demander_et_ecrire(parent, nom: str, format_: str, ecrire) -> None:
    """Demande où enregistrer (PDF ou PNG), écrit avec `ecrire(chemin)`, propose d'ouvrir le fichier."""
    if not autorise(parent):
        return
    types = {"pdf": ("PDF", "*.pdf"), "png": (t("Image PNG"), "*.png")}
    chemin = filedialog.asksaveasfilename(parent=parent, defaultextension="." + format_,
                                          initialfile=f"{nom_de_fichier(nom)} {date.today().isoformat()}.{format_}",
                                          filetypes=[types[format_]])
    if not chemin:
        return
    parent.configure(cursor="watch")
    parent.update_idletasks()
    try:
        ecrire(Path(chemin))
    except Exception as exc:
        messagebox.showerror(nom, str(exc), parent=parent)
        return
    finally:
        parent.configure(cursor="")
    if messagebox.askyesno(nom, t("Fichier enregistré : {fichier}.\n\nOuvrir le fichier ?", fichier=Path(chemin).name),
                           parent=parent):
        os.startfile(chemin)
