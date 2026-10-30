"""Rapprochement bancaire du compte Desjardins (1055) : relevés PDF de la banque <-> écritures de Sage.

Relevés : donnees/releves_desjardins/*.pdf (bouton « Charger des relevés… »), registre releves.json (décisions du
propriétaire par mouvement : compte Sage à utiliser, ignoré).
Lecture d'un relevé PDF numérique (pdfplumber) par colonnes : la ligne d'en-tête donne la position des colonnes
Retrait / Dépôt / Solde ; chaque montant est rangé dans la colonne la plus proche de son bord droit. Contrôle :
solde d'ouverture + dépôts − retraits = solde de fermeture.
Sage : lignes du compte 10550000 dans tjentact (montant > 0 = entrée) jointes à tjourent (date, source, description) ;
solde = taccount.dYtcLY (ouverture de l'exercice) + mouvements jusqu'à la date.
Rapprochement : même montant, même sens, date Sage de 7 jours avant à 3 jours après la date de banque (la plus
proche) ; sinon 2 ou 3 écritures Sage de la même période dont la somme égale le mouvement de banque.
Virements Interac reçus et dépôts d'espèces au guichet (règles du propriétaire, 2026-10-26) : ils font partie des
espèces des ventes ; absents de Sage, ils ne vont pas en écriture de journal : la page Encaissements les encaisse sur
les factures de vente (encaissements.especes_a_encaisser).
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from itertools import combinations
from pathlib import Path

from .factures_fournisseurs import MOIS

COMPTE_BANQUE = 10550000
AVANT, APRES = 7, 3  # fenêtre de dates acceptée pour une écriture Sage autour du mouvement de banque

RAPPROCHE = "Rapproché"
A_SAISIR = "Absent de Sage"
CONFIRME = "Compte choisi (à saisir)"
SAISI = "Saisi dans Sage"
IGNORE = "Ignoré"
EN_CIRCULATION = "Seulement dans Sage"
A_ENCAISSER = "À encaisser sur les ventes"

INTERAC = re.compile(r"virement\s+interac\s+de\b", re.I)
DEPOT_ESPECES = re.compile(r"^DGA\b|d.p.t au GA", re.I)  # « DGA Dépôt au GA » : espèces déposées au guichet
A_PAIE = "Paie absente de Sage"
SALAIRE = re.compile(r"salaire", re.I)  # « PPW Paiement internet à / Depot Salaire/PAY05012026 »
REFERENCE_PAIE = re.compile(r"PAYE?(\d{8})", re.I)
DELAI_PAIE = 120  # jours : un salaire peut être versé longtemps après la paie de Sage

MONTANT = re.compile(r"^-?\d{1,3}(?:[   ]?\d{3})*[.,]\d{2}-?$")


# --------------------------------------------------------------------------------------------- modèles
@dataclass
class MouvementBanque:
    date: str
    description: str
    montant: Decimal          # > 0 dépôt, < 0 retrait
    solde: Decimal | None = None
    releve: str = ""
    rang: int = 0             # ordre dans le relevé (deux mouvements identiques le même jour restent distincts)
    etat: str = A_SAISIR
    sage: list[int] = field(default_factory=list)  # n° d'écritures Sage rapprochées
    compte: str = ""          # compte Sage choisi pour la saisie (ex. 5450)
    remarque: str = ""

    @property
    def cle(self) -> str:
        return f"{self.releve}|{self.rang}"


@dataclass
class MouvementSage:
    ident: int
    date: str
    source: str
    description: str
    montant: Decimal
    rapproche: bool = False


@dataclass
class Releve:
    fichier: str
    debut: str = ""
    fin: str = ""
    solde_ouverture: Decimal | None = None
    solde_fermeture: Decimal | None = None
    mouvements: list[MouvementBanque] = field(default_factory=list)
    remarques: list[str] = field(default_factory=list)

    def controles(self) -> list[str]:
        problemes = list(self.remarques)
        if not self.mouvements:
            problemes.append("aucun mouvement lu")
        if self.solde_ouverture is not None and self.solde_fermeture is not None:
            calcule = self.solde_ouverture + sum((m.montant for m in self.mouvements), Decimal(0))
            if calcule != self.solde_fermeture:
                problemes.append(f"solde calculé {calcule} $ ≠ solde de fermeture {self.solde_fermeture} $")
        return problemes


# --------------------------------------------------------------------------------------------- lecture du PDF
def _montant(texte: str) -> Decimal:
    t = texte.replace(" ", "").replace(" ", "").replace(" ", "").replace("$", "")
    negatif = t.endswith("-") or t.startswith("-")
    t = t.strip("-").replace(",", ".") if t.count(",") == 1 and t.rfind(",") > t.rfind(".") else t.strip("-").replace(",", "")
    valeur = Decimal(t)
    return -valeur if negatif else valeur


def _date_ligne(texte: str, fin: date | None) -> date | None:
    """« 2026-09-02 », « 02/09 », « 02 SEP », « 2 sept. » ; l'année vient de la période du relevé."""
    t = texte.lower().replace(".", "").strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})$", t)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"(\d{1,2})[/ -]?([a-zéû]{3,9}|\d{1,2})$", t)
    if not m:
        return None
    jour = int(m[1])
    mois = int(m[2]) if m[2].isdigit() else MOIS.get(m[2])
    if not mois or not 1 <= mois <= 12:
        return None
    annee = fin.year if fin else date.today().year
    if fin and mois > fin.month + 1:  # relevé de janvier contenant des jours de décembre
        annee -= 1
    try:
        return date(annee, mois, jour)
    except ValueError:
        return None


def _lignes(page) -> list[list[dict]]:
    """Mots de la page regroupés en lignes (même hauteur à 3 points près), de gauche à droite."""
    lignes: list[list[dict]] = []
    for mot in sorted(page.extract_words(keep_blank_chars=False, x_tolerance=1.5), key=lambda m: (round(m["top"]), m["x0"])):
        if lignes and abs(lignes[-1][0]["top"] - mot["top"]) <= 3:
            lignes[-1].append(mot)
        else:
            lignes.append([mot])
    return [sorted(l, key=lambda m: m["x0"]) for l in lignes]


def _fusionner_nombres(mots: list[dict]) -> list[dict]:
    """« 1 » « 842,10 » (séparateur de milliers imprimé en espace) -> un seul mot « 1842,10 »."""
    sortie: list[dict] = []
    for mot in mots:
        precedent = sortie[-1] if sortie else None
        if (precedent and re.fullmatch(r"-?\d{1,3}", precedent["text"]) and re.match(r"\d{3}([ .,]|$)", mot["text"])
                and mot["x0"] - precedent["x1"] < 6):
            sortie[-1] = {**precedent, "text": precedent["text"] + mot["text"], "x1": mot["x1"]}
        else:
            sortie.append(dict(mot))
    return sortie


def _periode(texte: str) -> tuple[date | None, date | None]:
    """« du 1er février au 28 février 2026 » : la première date prend l'année de la seconde si elle n'en a pas."""
    from .factures_fournisseurs import lire_date

    m = re.search(r"\bdu\s+(\d{1,2})(?:er)?\s+([a-zéû]+)\.?\s*(\d{4})?\s+au\s+(\d{1,2})(?:er)?\s+([a-zéû]+)\.?\s+(\d{4})",
                  texte, re.IGNORECASE)
    if m:
        fin = lire_date(f"{m[4]} {m[5]} {m[6]}")
        annee = int(m[3]) if m[3] else (fin.year if fin else date.today().year)
        debut = lire_date(f"{m[1]} {m[2]} {annee}")
        if debut and fin and debut > fin:  # « du 1er décembre au 31 janvier 2027 »
            debut = debut.replace(year=debut.year - 1)
        return debut, fin
    dates = [d for d in (lire_date(x) for x in re.findall(r"\d{4}-\d{2}-\d{2}", texte)) if d]
    return (min(dates), max(dates)) if dates else (None, None)


COLONNES = {"frais": "frais", "retrait": "retrait", "retraits": "retrait", "débit": "retrait", "débits": "retrait",
            "debit": "retrait", "dépôt": "depot", "dépôts": "depot", "depot": "depot", "depots": "depot",
            "crédit": "depot", "crédits": "depot", "credit": "depot", "solde": "solde"}
# Fin du compte principal : les sections suivantes du relevé (parts sociales, sommaire des frais) sont ignorées
FIN_DU_COMPTE = re.compile(r"^(compte d'[ée]pargne et de placement|sommaire des frais|part de qualification)")


_RELEVES_LUS: dict[str, tuple] = {}


def _releve_en_dict(r: Releve) -> dict:
    from dataclasses import asdict

    return json.loads(json.dumps(asdict(r), default=str))


def _releve_depuis_dict(d: dict) -> Releve:
    def dec(v):
        return None if v is None else Decimal(v)

    mouvements = [MouvementBanque(m["date"], m["description"], Decimal(m["montant"]), dec(m["solde"]), m["releve"],
                                  m["rang"], m["etat"], list(m["sage"]), m["compte"], m["remarque"])
                  for m in d["mouvements"]]
    return Releve(d["fichier"], d["debut"], d["fin"], dec(d["solde_ouverture"]), dec(d["solde_fermeture"]),
                  mouvements, list(d["remarques"]))


def lire_releve_garde(chemin: Path) -> Releve:
    """Relevé lu une seule fois tant que le fichier ne change pas (date et taille) : gardé en mémoire et dans
    « releves_lus.cache.json » du dossier des relevés. Lire les PDF prenait environ 3 s à chaque suivi des ventes
    ou rapprochement (mesuré le 2026-10-30). Chaque appel rend une copie neuve (les décisions s'y appliquent)."""
    chemin = Path(chemin)
    etat = chemin.stat()
    signature = [etat.st_mtime_ns, etat.st_size]
    deja = _RELEVES_LUS.get(str(chemin))
    if deja and deja[0] == signature:
        return _releve_depuis_dict(deja[1])
    cache = chemin.parent / "releves_lus.cache.json"
    try:
        tous = json.loads(cache.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        tous = {}
    entree = tous.get(chemin.name)
    if entree and entree.get("signature") == signature:
        _RELEVES_LUS[str(chemin)] = (signature, entree["releve"])
        return _releve_depuis_dict(entree["releve"])
    releve = lire_releve(chemin)
    donnees = _releve_en_dict(releve)
    _RELEVES_LUS[str(chemin)] = (signature, donnees)
    tous[chemin.name] = {"signature": signature, "releve": donnees}
    try:
        cache.write_text(json.dumps(tous, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass
    return _releve_depuis_dict(donnees)


def lire_releve(chemin: Path) -> Releve:
    """Relevé Desjardins : colonnes Date, Code, Description, Frais, Retrait, Dépôt, Solde ; « 6 840.89 » ;
    description parfois sur deux lignes (les montants sont alors sur la seconde)."""
    import pdfplumber

    releve = Releve(Path(chemin).name)
    with pdfplumber.open(chemin) as pdf:
        texte = "\n".join((p.extract_text() or "") for p in pdf.pages)
        debut, fin = _periode(texte)
        releve.debut, releve.fin = (debut.isoformat() if debut else ""), (fin.isoformat() if fin else "")
        colonnes: dict[str, float] = {}
        solde_courant: Decimal | None = None
        en_attente: tuple[date, str] | None = None  # ligne datée dont les montants sont sur la ligne suivante
        termine = False
        for page in pdf.pages:
            if termine:
                break
            for mots in _lignes(page.dedupe_chars()):
                mots = _fusionner_nombres(mots)
                bas = " ".join(m["text"] for m in mots).lower()
                if colonnes and FIN_DU_COMPTE.match(bas):
                    termine = True
                    break
                if "solde" in bas and "description" in bas:  # en-tête des colonnes (répété sur chaque page)
                    colonnes = {COLONNES[m["text"].lower().strip(".:")]: m["x1"] for m in mots
                                if m["text"].lower().strip(".:") in COLONNES}
                    continue
                montants = [m for m in mots if MONTANT.match(m["text"])]
                autres = [m for m in mots if m not in montants]
                if re.search(r"solde (reporté|précédent|d'ouverture|au début|initial)|ancien solde", bas) and montants:
                    if releve.solde_ouverture is None:
                        releve.solde_ouverture = solde_courant = _montant(montants[-1]["text"])
                    continue
                if re.search(r"(nouveau solde|solde de fermeture|solde à la fin|solde final)", bas) and montants:
                    releve.solde_fermeture = _montant(montants[-1]["text"])
                    continue
                if not colonnes:
                    continue
                jour = None
                for n in (2, 1):  # « 2 FEV » ou « 2026-09-02 »
                    if len(autres) >= n:
                        jour = _date_ligne(" ".join(m["text"] for m in autres[:n]), fin)
                        if jour:
                            autres = autres[n:]
                            break
                description = " ".join(m["text"] for m in autres).strip()
                if not montants:
                    if jour and description:
                        en_attente = (jour, description)
                    elif en_attente and description:
                        en_attente = (en_attente[0], f"{en_attente[1]} {description}")
                    elif releve.mouvements and description and not jour:
                        releve.mouvements[-1].description += f" {description}"
                    continue
                if not jour and en_attente:  # suite de la ligne précédente, avec les montants
                    jour, description = en_attente[0], f"{en_attente[1]} {description}".strip()
                en_attente = None
                valeurs = {"frais": None, "retrait": None, "depot": None, "solde": None}
                for m in montants:
                    colonne = min(colonnes, key=lambda c: abs(colonnes[c] - m["x1"]))
                    valeurs[colonne] = _montant(m["text"])
                if all(valeurs[c] is None for c in ("frais", "retrait", "depot")):
                    continue
                montant = (valeurs["depot"] or Decimal(0)) - (valeurs["retrait"] or Decimal(0)) - (
                    valeurs["frais"] or Decimal(0))
                jour = jour or (date.fromisoformat(releve.mouvements[-1].date) if releve.mouvements else None)
                if not jour:
                    releve.remarques.append(f"ligne sans date ignorée : {bas}")
                    continue
                mouvement = MouvementBanque(jour.isoformat(), description, montant, valeurs["solde"], releve.fichier,
                                            len(releve.mouvements))
                if valeurs["solde"] is not None and solde_courant is not None and solde_courant + montant != valeurs["solde"]:
                    mouvement.remarque = f"solde attendu {solde_courant + montant} $, imprimé {valeurs['solde']} $"
                solde_courant = valeurs["solde"] if valeurs["solde"] is not None else (
                    None if solde_courant is None else solde_courant + montant)
                releve.mouvements.append(mouvement)
        if not colonnes:
            releve.remarques.append("colonnes Retrait / Dépôt / Solde introuvables (relevé numérisé ?)")
        if releve.solde_fermeture is None and solde_courant is not None and releve.mouvements:
            releve.solde_fermeture = solde_courant
    return releve


# --------------------------------------------------------------------------------------------- Sage
def mouvements_sage(base, debut: date, fin: date, compte: int = COMPTE_BANQUE) -> list[MouvementSage]:
    """Écritures du compte, sans les écritures corrigées ni leurs contreparties (tjourent.nType 64/65 : original
    corrigé, 32/33 : contrepassation « CORR … Contrep. de Jn ») : elles s'annulent et ne passent pas à la banque."""
    lignes = base.requete(
        "SELECT e.lId, e.dtJourDate, e.sSource, e.sComment, ROUND(SUM(a.dAmount), 2) montant FROM tjentact a "
        "JOIN tjourent e ON e.lId = a.lJEntId WHERE a.lAcctId = %s AND e.dtJourDate BETWEEN %s AND %s "
        "AND e.nType NOT IN (32, 33, 64, 65) "
        "GROUP BY e.lId, e.dtJourDate, e.sSource, e.sComment HAVING ROUND(SUM(a.dAmount), 2) <> 0 "
        "ORDER BY e.dtJourDate, e.lId", (compte, debut.isoformat(), fin.isoformat()))
    return [MouvementSage(l["lId"], l["dtJourDate"].date().isoformat(), (l["sSource"] or "").strip(),
                          (l["sComment"] or "").strip(), Decimal(str(l["montant"]))) for l in lignes]


def solde_sage(base, jour: date, compte: int = COMPTE_BANQUE) -> Decimal:
    """Solde du compte à la fin de `jour` : ouverture de l'exercice + écritures jusqu'à ce jour."""
    ouverture = base.requete("SELECT dYtcLY FROM taccount WHERE lId = %s", (compte,))
    mouvements = base.requete("SELECT ROUND(SUM(a.dAmount), 2) s FROM tjentact a JOIN tjourent e ON e.lId = a.lJEntId "
                              "WHERE a.lAcctId = %s AND e.dtJourDate <= %s", (compte, jour.isoformat()))
    return (Decimal(str(ouverture[0]["dYtcLY"] if ouverture else 0)) + Decimal(str(mouvements[0]["s"] or 0))
            ).quantize(Decimal("0.01"))


def factures_ventes(base, client: str) -> dict[str, date]:
    """Factures de vente du client « Vente Hebdomadaire » dans Sage : n° -> date de la facture."""
    return {l["s"]: l["d"].date() if hasattr(l["d"], "date") else l["d"] for l in base.requete(
        "SELECT TRIM(t.sSource) s, t.dtDate d FROM tcustr t JOIN tcustomr c ON c.lId = t.lCusId "
        "WHERE c.sName = %s AND t.sSource LIKE 'Vente %%'", (client,))}


def comptes_sage(base) -> list[str]:
    """« 5450 Frais bancaires » : comptes actifs, pour le choix du compte d'une écriture à saisir."""
    from .base_sage import numero_compte

    return [f"{numero_compte(l['lId'])} {l['sName'].strip()}"
            for l in base.requete("SELECT lId, sName FROM taccount WHERE bInactive = 0 ORDER BY lId")]


# --------------------------------------------------------------------------------------------- rapprochement
def _dans_fenetre(banque: MouvementBanque, sage: MouvementSage) -> bool:
    ecart = (date.fromisoformat(banque.date) - date.fromisoformat(sage.date)).days
    return -APRES <= ecart <= AVANT


def _date_reference(texte: str) -> date | None:
    """« PAY05012026 », « PAYE05012026 » (jjmmaaaa) ou « PAY20260119 » (aaaammjj) -> date de la paie."""
    m = REFERENCE_PAIE.search(texte or "")
    if not m:
        return None
    chiffres = m.group(1)
    for jour, mois, annee in ((chiffres[:2], chiffres[2:4], chiffres[4:]), (chiffres[6:], chiffres[4:6], chiffres[:4])):
        try:
            return date(int(annee), int(mois), int(jour))
        except ValueError:
            continue
    return None


def est_salaire(m: MouvementBanque) -> bool:
    return m.montant < 0 and bool(SALAIRE.search(m.description))


def _rapprocher_salaires(a_relier: list[MouvementBanque], sage: list[MouvementSage],
                         anterieures: list[MouvementSage]) -> None:
    """Dépôts de salaire (règle du propriétaire) : reliés aux paies de Sage (source « PAY… ») du même montant, même
    payées plusieurs semaines plus tard ; la paie dont la référence (date) est celle du relevé passe en premier."""
    for b in sorted((x for x in a_relier if est_salaire(x)), key=lambda x: x.date):
        jour, reference = date.fromisoformat(b.date), _date_reference(b.description)
        candidates = [s for s in sage + anterieures if not s.rapproche and s.montant == b.montant
                      and s.source.upper().startswith("PAY")
                      and -APRES <= (jour - date.fromisoformat(s.date)).days <= DELAI_PAIE]
        if not candidates:
            continue
        s = min(candidates, key=lambda c: (reference is None or _date_reference(c.source) != reference,
                                           abs((jour - date.fromisoformat(c.date)).days)))
        s.rapproche, b.etat, b.sage = True, RAPPROCHE, [s.ident]
        if s not in sage:
            sage.append(s)


def rapprocher(banque: list[MouvementBanque], sage: list[MouvementSage],
               anterieures: list[MouvementSage] | None = None) -> None:
    """Relie les mouvements de banque aux écritures Sage (ne touche pas aux mouvements saisis, choisis ou ignorés).
    `anterieures` : paies de Sage d'avant la période des relevés (salaires versés en retard) ; celles qui sont reliées
    sont ajoutées à `sage`."""
    for s in sage:
        s.rapproche = False
    a_relier = [b for b in banque if b.etat in (A_SAISIR, RAPPROCHE, A_ENCAISSER, A_PAIE)]
    for b in a_relier:
        b.etat, b.sage = A_SAISIR, []
    for b in banque:  # mouvements déjà saisis par l'application : leurs écritures restent prises
        if b.etat == SAISI:
            for s in sage:
                if s.ident in b.sage:
                    s.rapproche = True
    # 0. dépôts de salaire -> paies de Sage
    _rapprocher_salaires(a_relier, sage, [s for s in anterieures or [] if s.ident not in {x.ident for x in sage}])
    # 1. une écriture du même montant, à la date la plus proche
    for b in sorted(a_relier, key=lambda x: x.date):
        candidates = [s for s in sage if not s.rapproche and s.montant == b.montant and _dans_fenetre(b, s)]
        if candidates:
            s = min(candidates, key=lambda c: abs((date.fromisoformat(b.date) - date.fromisoformat(c.date)).days))
            s.rapproche, b.etat, b.sage = True, RAPPROCHE, [s.ident]
    # 2. deux ou trois écritures dont la somme égale le mouvement (ex. dépôt de fin de semaine)
    for b in [x for x in a_relier if x.etat == A_SAISIR]:
        libres = [s for s in sage if not s.rapproche and _dans_fenetre(b, s) and (s.montant > 0) == (b.montant > 0)][:14]
        for taille in (2, 3):
            groupe = next((g for g in combinations(libres, taille) if sum(s.montant for s in g) == b.montant), None)
            if groupe:
                for s in groupe:
                    s.rapproche = True
                b.etat, b.sage = RAPPROCHE, [s.ident for s in groupe]
                break
    # 3. Interac reçu ou dépôt d'espèces absent de Sage : à encaisser sur les factures de vente (page Encaissements)
    for b in a_relier:
        if b.etat == A_SAISIR and est_especes(b):
            b.etat = A_ENCAISSER
        elif b.etat == A_SAISIR and est_salaire(b):
            b.etat = A_PAIE  # la paie se saisit dans le module Paie de Sage, pas en écriture de journal


def est_interac(m: MouvementBanque) -> bool:
    return m.montant > 0 and bool(INTERAC.search(m.description))


def est_especes(m: MouvementBanque) -> bool:
    """Argent des ventes encaissé hors Clover : virement Interac reçu ou dépôt d'espèces au guichet automatique."""
    return m.montant > 0 and (est_interac(m) or bool(DEPOT_ESPECES.search(m.description)))


def jours_couverts(numero: str) -> list[date]:
    """Jours de vente d'une facture : jour, fin de semaine, « Vente S43-2026 » ou « Vente Octobre 2026 »."""
    from .facture import MOIS as MOIS_FACTURE
    from .facture import jours_facture

    jours = jours_facture(numero)
    if jours:
        return jours
    m = re.fullmatch(r"Vente S(\d{2})-(\d{4})", numero.strip())
    if m:
        lundi = date.fromisocalendar(int(m.group(2)), int(m.group(1)), 1)
        return [lundi + timedelta(days=i) for i in range(7)]
    m = re.fullmatch(r"Vente (\S+) (\d{4})", numero.strip())
    if m and m.group(1) in MOIS_FACTURE:
        premier = date(int(m.group(2)), MOIS_FACTURE.index(m.group(1)) + 1, 1)
        return [premier + timedelta(days=i) for i in range(31) if (premier + timedelta(days=i)).month == premier.month]
    return []


def facture_du_jour(jour: date, factures: dict[str, date] | set[str]) -> str:
    """Facture de vente qui couvre `jour` (la plus précise : jour, puis fin de semaine, semaine, mois) ; « » sinon."""
    candidates = [(len(jours), n) for n in factures if jour in (jours := jours_couverts(n))]
    return min(candidates)[1] if candidates else ""


def resume(banque: list[MouvementBanque], sage: list[MouvementSage], solde_banque: Decimal | None,
           solde_de_sage: Decimal | None, ouverture_banque: Decimal | None = None,
           ouverture_sage: Decimal | None = None, debut: date | None = None) -> dict:
    """Chiffres du rapprochement. Écart de fin = écart déjà présent à l'ouverture de la période + mouvements de banque
    absents de Sage − écritures Sage pas encore passées à la banque ; le reste est « non expliqué ».

    Un mouvement relié à une écriture Sage d'AVANT le `debut` de la période (salaire versé en retard, écriture saisie
    quelques jours plus tôt) compte comme absent : l'écriture est déjà dans l'écart d'ouverture."""
    dates = {s.ident: s.date for s in sage}
    avant = (lambda b: debut is not None and b.etat == RAPPROCHE and b.sage
             and all(n in dates and date.fromisoformat(dates[n]) < debut for n in b.sage))
    absents = sum((b.montant for b in banque if b.etat in (A_SAISIR, CONFIRME, A_ENCAISSER, A_PAIE) or avant(b)),
                  Decimal(0))
    circulation = sum((s.montant for s in sage if not s.rapproche), Decimal(0))
    ecart = None if solde_banque is None or solde_de_sage is None else solde_banque - solde_de_sage
    ouverture = (None if ouverture_banque is None or ouverture_sage is None else ouverture_banque - ouverture_sage)
    explique = absents - circulation + (ouverture or Decimal(0))
    return {"solde_banque": solde_banque, "solde_sage": solde_de_sage, "ecart": ecart, "ecart_ouverture": ouverture,
            "absents": absents, "circulation": circulation, "non_explique": None if ecart is None else ecart - explique}


# --------------------------------------------------------------------------------------------- registre
class RegistreReleves:
    """Relevés chargés et décisions du propriétaire (donnees/releves_desjardins/releves.json)."""

    def __init__(self, racine: Path):
        self.dossier = Path(racine) / "donnees" / "releves_desjardins"
        self.chemin = self.dossier / "releves.json"
        self.decisions: dict[str, dict] = {}   # clé de mouvement -> {"etat", "compte", "sage", "remarque"}
        self.memoire: dict[str, str] = {}      # description normalisée -> compte Sage (pour les mouvements semblables)
        if self.chemin.exists():
            donnees = json.loads(self.chemin.read_text(encoding="utf-8"))
            self.decisions, self.memoire = donnees.get("decisions", {}), donnees.get("memoire", {})

    def enregistrer(self) -> None:
        self.dossier.mkdir(parents=True, exist_ok=True)
        provisoire = self.chemin.with_suffix(".tmp")
        provisoire.write_text(json.dumps({"decisions": self.decisions, "memoire": self.memoire}, ensure_ascii=False,
                                         indent=1), encoding="utf-8")
        provisoire.replace(self.chemin)

    def fichiers(self) -> list[Path]:
        return sorted(p for p in self.dossier.glob("*.pdf")) if self.dossier.exists() else []

    def lire(self) -> list[Releve]:
        releves = []
        for chemin in self.fichiers():
            try:
                r = lire_releve_garde(chemin)
            except Exception as exc:
                r = Releve(chemin.name, remarques=[f"lecture impossible : {exc}"])
            for m in r.mouvements:
                d = self.decisions.get(m.cle)
                if d:
                    m.etat, m.compte, m.sage, m.remarque = d["etat"], d.get("compte", ""), d.get("sage", []), \
                        d.get("remarque", m.remarque)
                elif not m.compte:
                    m.compte = self.memoire.get(motif(m.description), "")
            releves.append(r)
        return releves

    def retenir(self, m: MouvementBanque) -> None:
        if m.etat in (CONFIRME, SAISI, IGNORE):
            self.decisions[m.cle] = {"etat": m.etat, "compte": m.compte, "sage": m.sage, "remarque": m.remarque}
        else:
            self.decisions.pop(m.cle, None)
        if m.compte and m.etat in (CONFIRME, SAISI):
            self.memoire[motif(m.description)] = m.compte

    def supprimer(self, fichier: str) -> None:
        """Retire un relevé : le fichier va dans « supprimes », ses décisions sont oubliées."""
        chemin = self.dossier / fichier
        if chemin.exists():
            corbeille = self.dossier / "supprimes"
            corbeille.mkdir(exist_ok=True)
            chemin.replace(corbeille / chemin.name)
        self.decisions = {k: v for k, v in self.decisions.items() if not k.startswith(f"{fichier}|")}


def motif(description: str) -> str:
    """Description sans chiffres ni ponctuation : « FRAIS MENS. 09/2026 » et « FRAIS MENS. 10/2026 » se ressemblent."""
    from .correspondance import normaliser

    return " ".join(re.sub(r"[^a-z ]", " ", normaliser(description)).split())


# --------------------------------------------------------------------------------------------- rapport
def ecrire_rapport(chemin: Path, releves: list[Releve], banque: list[MouvementBanque], sage: list[MouvementSage],
                   chiffres: dict, periode: tuple[date, date]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    vert = PatternFill("solid", fgColor="1F7A3A")
    blanc_gras = Font(bold=True, color="FFFFFF")
    classeur = Workbook()
    f = classeur.active
    f.title = "Rapprochement"
    f.append([f"Rapprochement bancaire — compte 1055 Desjardins — du {periode[0]} au {periode[1]}"])
    f["A1"].font = Font(bold=True, size=14, color="13772D")
    f.append([])
    for libelle, cle in (("Solde du relevé (banque)", "solde_banque"), ("Solde du compte 1055 dans Sage", "solde_sage"),
                         ("Écart (banque − Sage)", "ecart"), ("Écart déjà présent à l'ouverture", "ecart_ouverture"),
                         ("Mouvements de banque absents de Sage", "absents"),
                         ("Écritures Sage pas encore à la banque", "circulation"), ("Écart non expliqué", "non_explique")):
        f.append([libelle, float(chiffres[cle]) if chiffres[cle] is not None else ""])
    f.append([])
    f.append(["Relevés", "Période", "Solde d'ouverture", "Solde de fermeture", "Contrôle"])
    for r in releves:
        f.append([r.fichier, f"{r.debut} au {r.fin}", float(r.solde_ouverture or 0), float(r.solde_fermeture or 0),
                  "; ".join(r.controles()) or "OK"])
    f.column_dimensions["A"].width = 46
    for col in "BCDE":
        f.column_dimensions[col].width = 22
    for titre, entetes, lignes in (
            ("Banque", ["Date", "Description", "Montant", "Solde", "État", "Écritures Sage", "Compte", "Relevé"],
             [[b.date, b.description, float(b.montant), float(b.solde) if b.solde is not None else "", b.etat,
               ", ".join(map(str, b.sage)), b.compte, b.releve] for b in banque]),
            ("Sage", ["N° écriture", "Date", "Source", "Description", "Montant", "Rapprochée"],
             [[s.ident, s.date, s.source, s.description, float(s.montant), "oui" if s.rapproche else "NON"]
              for s in sage])):
        feuille = classeur.create_sheet(titre)
        feuille.append(entetes)
        for cellule in feuille[1]:
            cellule.fill, cellule.font = vert, blanc_gras
        for ligne in lignes:
            feuille.append(ligne)
        for i, largeur in enumerate((12, 46, 14, 14, 24, 16, 12, 30)[:len(entetes)]):
            feuille.column_dimensions[chr(65 + i)].width = largeur
        feuille.freeze_panes = "A2"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin


def restreindre(releves: list[Releve], du: date, au: date):
    """Mouvements de banque du `du` au `au` inclus, solde à l'ouverture du `du` et solde à la fin du `au`.

    Solde courant suivi relevé par relevé (dans l'ordre des dates) : solde d'ouverture du relevé, puis chaque
    mouvement ; le solde imprimé sur la ligne, quand il existe, fait foi."""
    banque, ouverture, fermeture = [], None, None
    for r in sorted(releves, key=lambda r: r.debut or "9999"):
        courant = r.solde_ouverture
        if r.debut and date.fromisoformat(r.debut) >= du and ouverture is None:
            ouverture = courant
        for m in r.mouvements:
            jour = date.fromisoformat(m.date)
            if jour < du:
                pass
            elif jour <= au:
                if ouverture is None:
                    ouverture = courant
                banque.append(m)
            courant = m.solde if m.solde is not None else (None if courant is None else courant + m.montant)
            if jour <= au:
                fermeture = courant
        if r.fin and date.fromisoformat(r.fin) <= au and r.solde_fermeture is not None:
            fermeture = r.solde_fermeture
    return banque, ouverture, fermeture


def periode_des(releves: list[Releve]) -> tuple[date, date] | None:
    jours = [date.fromisoformat(m.date) for r in releves for m in r.mouvements]
    debuts = [date.fromisoformat(r.debut) for r in releves if r.debut]
    fins = [date.fromisoformat(r.fin) for r in releves if r.fin]
    if not jours and not debuts:
        return None
    return min(debuts + jours), max(fins + jours)

