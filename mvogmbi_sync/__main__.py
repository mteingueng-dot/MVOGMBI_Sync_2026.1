"""Ligne de commande.

Exemples :
    python -m mvogmbi_sync rapport --periode jour --date 2026-01-15
    python -m mvogmbi_sync rapport --periode semaine --date 2026-01-15 --par-jour
    python -m mvogmbi_sync rapport --periode mois --date 2026-01-01 --export "Ventes Janvier.xlsx"
    python -m mvogmbi_sync verifier
    python -m mvogmbi_sync comparer --date 2026-10-23 --export "Ventes 23 octobre.xlsx"
"""

from __future__ import annotations

import argparse
import sys
import tomllib
from datetime import date, timedelta
from pathlib import Path

from .models import RapportVentes
from .periodes import TypePeriode, blocs_facturation, bornes, jours
from .sortie_excel import ecrire_rapports

import os

# Dossier de l'application : celui du programme installé (.exe), sinon celui du projet ;
# MVOGMBI_RACINE permet de le choisir explicitement.
if os.environ.get("MVOGMBI_RACINE"):
    RACINE = Path(os.environ["MVOGMBI_RACINE"])
elif getattr(sys, "frozen", False):
    RACINE = Path(sys.executable).resolve().parent
else:
    RACINE = Path(__file__).resolve().parent.parent


NOM_EXE_CONSOLE = "MVOGMBI_Sync_console.exe"


def commande_moteur() -> list[str]:
    """Commande qui lance le moteur en ligne de commande (programme installé ou Python)."""
    if getattr(sys, "frozen", False):
        return [str(Path(sys.executable).resolve().with_name(NOM_EXE_CONSOLE))]
    return [sys.executable, "-m", "mvogmbi_sync"]


def charger_config(chemin: Path):
    from .api_clover import ConfigClover

    if not chemin.exists():
        sys.exit(f"Fichier de configuration introuvable : {chemin}\n"
                 "Copiez config.example.toml en config.toml et complétez-le.")
    from . import coffre

    donnees = tomllib.loads(chemin.read_text(encoding="utf-8"))
    clover = donnees.get("clover", {})
    jeton = coffre.lire(coffre.CLOVER_JETON, clover.get("api_token"))
    if not clover.get("merchant_id") or not jeton:
        sys.exit("Identifiant marchand ou jeton Clover manquant : complétez la configuration "
                 "(application MVOGMBI_Sync, onglet Clover).")
    return ConfigClover(
        merchant_id=clover.get("merchant_id", ""),
        api_token=jeton,
        base_url=clover.get("base_url", "https://api.clover.com"),
        fuseau_horaire=donnees.get("general", {}).get("fuseau_horaire", "America/Toronto"),
    )


def afficher(rapport: RapportVentes) -> None:
    print(f"\n{rapport.debut} → {rapport.fin}  ({rapport.source})")
    print(f"  Articles distincts : {len(rapport.lignes)}")
    print(f"  Ventes brutes      : {rapport.total_brut:>12,.2f} $")
    print(f"  Rabais             : {rapport.total_rabais:>12,.2f} $")
    print(f"  Remboursements     : {rapport.total_remboursements:>12,.2f} $")
    print(f"  Ventes nettes      : {rapport.total_net:>12,.2f} $")


def commande_encaissements(args: argparse.Namespace) -> None:
    """Dépôts Clover (lots de cartes) à encaisser sur les factures journalières, et leur état dans Sage."""
    from .api_clover import ClientClover
    from .encaissements import calculer, ecrire_excel, numero_compte_sage
    from .facture import CLIENT

    config_toml = tomllib.loads(Path(args.config).read_text(encoding="utf-8"))
    compte = str(config_toml.get("encaissements", {}).get("compte_depot", "1055"))
    client = ClientClover(charger_config(Path(args.config)))
    depots = lire_sage(config_toml, lambda base: calculer(client, base, CLIENT, numero_compte_sage(compte),
                                                          date.fromisoformat(args.du), date.fromisoformat(args.au)),
                       config_toml.get("sage_import", {}).get("fichier_entreprise"))
    for d in depots:
        print(f"{d.date} ({d.date:%a})  dépôt {d.montant:>10,.2f} $  au compte {compte}  : {d.etat}")
        for l in d.lignes:
            print(f"    {l.facture:<18} {l.montant:>10,.2f} $   dû dans Sage : "
                  f"{'' if l.montant_du is None else f'{l.montant_du:,.2f} $':>12}   {l.etat}")
    if args.export:
        ecrire_excel(depots, Path(args.export), compte)
        print(f"Liste enregistrée : {args.export}")
    if args.appliquer:
        from .robot_encaissements import appliquer

        mode = str(config_toml.get("encaissements", {}).get("mode", "Comptant"))
        resultats = appliquer(config_toml, depots, CLIENT, compte, mode)
        if any(not r.traite for r in resultats):
            sys.exit(1)


def commande_encaisser_bureau(args: argparse.Namespace) -> None:
    """Commande interne, lancée sur le bureau caché par le traitement automatique ou la page Encaissements."""
    import json

    from .encaissements import encaisser_sur_bureau_cache

    config_toml = tomllib.loads(Path(args.config).read_text(encoding="utf-8"))
    notes = []
    bilan = encaisser_sur_bureau_cache(RACINE, config_toml, journal=lambda t: notes.append(t))
    bilan["notes"] = notes
    Path(args.resultat).write_text(json.dumps(bilan, ensure_ascii=False, indent=1), encoding="utf-8")
    sys.exit(1 if bilan["erreur"] or bilan["a_verifier"] else 0)


def commande_rapport(args: argparse.Namespace) -> None:
    debut, fin = bornes(TypePeriode(args.periode), date.fromisoformat(args.date))

    if args.export:
        if args.par_jour:
            sys.exit("--par-jour n'est pas possible avec un export : le fichier Clover ne contient pas les dates.")
        from .export_clover import lire_export
        rapports = [lire_export(Path(args.export), debut, fin)]
    else:
        from .api_clover import ClientClover
        client = ClientClover(charger_config(Path(args.config)))
        periodes = [(j, j) for j in jours(debut, fin)] if args.par_jour else [(debut, fin)]
        rapports = [client.rapport(d, f) for d, f in periodes]

    for rapport in rapports:
        afficher(rapport)
    sortie = RACINE / "sortie" / f"ventes_{args.periode}_{debut}_{fin}.xlsx"
    print(f"\nRapport Excel : {ecrire_rapports(rapports, sortie)}")


def commande_verifier(args: argparse.Namespace) -> None:
    import requests

    from .api_clover import ClientClover

    client = ClientClover(charger_config(Path(args.config)))
    try:
        infos = client.marchand()
    except requests.HTTPError as exc:
        code = exc.response.status_code
        raison = {401: "jeton refusé (invalide ou expiré)",
                  403: "droits insuffisants sur le jeton",
                  404: "identifiant marchand introuvable"}.get(code, "erreur inattendue")
        sys.exit(f"Échec de la connexion à Clover : {code} — {raison}.")
    print(f"Connexion à Clover réussie : {infos.get('name', '(nom inconnu)')}")
    hier = date.today() - timedelta(days=1)
    nombre = sum(1 for _ in client.commandes(hier, hier))
    print(f"Commandes trouvées pour hier ({hier}) : {nombre}")


def commande_comparer(args: argparse.Namespace) -> None:
    from .api_clover import ClientClover
    from .comparaison import comparer, ecrire_comparaison
    from .export_clover import lire_export

    jour = date.fromisoformat(args.date)
    export = lire_export(Path(args.export), jour, jour)
    api = ClientClover(charger_config(Path(args.config))).rapport(jour, jour)
    afficher(export)
    afficher(api)
    ecarts = comparer(export, api)
    sortie = ecrire_comparaison(ecarts, RACINE / "sortie" / f"comparaison_{jour}.xlsx")
    if ecarts:
        print(f"\n{len(ecarts)} article(s) en écart. Détail : {sortie}")
    else:
        print(f"\nAucun écart : l'API et l'export concordent au cent près. Détail : {sortie}")


def ouvrir_base_sage(config_toml: dict, fichier_sai: str | Path | None = None):
    """Base Sage du fichier d'entreprise indiqué (sinon celui de [sage_base], sinon le port configuré)."""
    from . import coffre
    from .base_sage import BaseSage, ConfigSageBase, port_du_fichier

    b = config_toml.get("sage_base", {})
    config = ConfigSageBase(
        utilisateur=b.get("utilisateur", ""), mot_de_passe=coffre.lire(coffre.SAGE_LECTURE, b.get("mot_de_passe")),
        hote=b.get("hote", "127.0.0.1"), port=int(b.get("port", 13540)), base=b.get("base", "simply"),
    )
    fichier = fichier_sai or b.get("fichier_entreprise")
    if fichier:
        config.port = port_du_fichier(Path(fichier), config)
    return BaseSage(config)


COPIE_ARTICLES_SAGE = "articles_sage.csv"


def lire_sage(config_toml: dict, fonction, fichier_sai=None):
    """Exécute fonction(base) sur la base Sage du fichier ; si Sage est fermé, la base est ouverte par l'outil
    d'import le temps de la lecture, puis refermée (aucun import)."""
    from .sage_auto import lire_pendant_connexion

    fichier = fichier_sai or config_toml.get("sage_base", {}).get("fichier_entreprise") \
        or config_toml.get("sage_import", {}).get("fichier_entreprise")

    def lecture():
        with ouvrir_base_sage(config_toml, fichier) as base:
            return fonction(base)

    try:
        return lecture()
    except Exception:
        exe = executable_import()
        if not fichier or not exe:
            raise
        return lire_pendant_connexion(exe, Path(fichier), lecture)


def commande_verifier_sage(args: argparse.Namespace) -> None:
    config_toml = tomllib.loads(Path(args.config).read_text(encoding="utf-8")) if Path(args.config).exists() else {}
    if not config_toml.get("sage_base", {}).get("utilisateur"):
        sys.exit("Renseignez la section [sage_base] (utilisateur, mot_de_passe) dans config.toml.")
    try:
        with ouvrir_base_sage(config_toml) as base:
            table, cols = base.table_articles()
            articles = base.articles()
            print(f"Connexion à Sage réussie. Table des articles : {table} ({len(articles)} articles)")
            print("Colonnes :", ", ".join(cols.values()))
            for a in articles[:5]:
                print(f"   {a.numero:<16} {a.description}  (compte {a.compte_revenus})")
    except Exception as exc:
        sys.exit(f"Connexion à Sage impossible : {exc}\n"
                 "Vérifiez que Sage 50 est ouvert sur le fichier de l'entreprise et que l'utilisateur "
                 "a le droit « Accès par des tiers ».")


def verifier_comptes_sage(config_toml: dict, fichier_sai=None, numeros: set[str] | None = None,
                          fichier_sortie: Path | None = None) -> list:
    """Vérifie la règle 15xx -> 40xx / 50xx ; affiche un résumé et écrit le détail dans un fichier Excel."""
    from .verification_comptes import ecrire_anomalies, verifier

    with ouvrir_base_sage(config_toml, fichier_sai) as base:
        articles = base.comptes_articles()
        noms = base.noms_comptes()
        invisibles = base.numeros_avec_caracteres_invisibles()
    for numero, description in invisibles.items():
        if numeros is None or numero in numeros:
            print(f"  ! Numéro d'article avec caractères invisibles (à retaper dans Sage) : {numero} {description}")
    if numeros is not None:
        articles = [a for a in articles if a.numero in numeros]
    anomalies = verifier(articles, noms)
    bloquantes = sum(1 for x in anomalies if x.bloquante)
    print(f"Comptes des articles Sage : {len(articles)} vérifiés, {bloquantes} sans compte de produits (bloquant), "
          f"{len(anomalies) - bloquantes} incohérent(s) avec la règle 15xx -> 40xx / 50xx")
    if anomalies and fichier_sortie:
        print(f"  Détail : {ecrire_anomalies(anomalies, fichier_sortie)}")
    return anomalies


def rapport_comptes(fichier_sai) -> Path:
    """Rapport des comptes à corriger, un par fichier d'entreprise (vraie comptabilité, copie d'essai…)."""
    suffixe = f"_{Path(fichier_sai).stem}" if fichier_sai else ""
    return RACINE / "donnees" / f"comptes_a_corriger{suffixe}.xlsx"


def commande_verifier_comptes(args: argparse.Namespace) -> None:
    config_toml = lire_toml(args.config)
    fichier = args.fichier or config_toml.get("sage_base", {}).get("fichier_entreprise")
    try:
        anomalies = verifier_comptes_sage(config_toml, fichier, fichier_sortie=rapport_comptes(fichier))
    except Exception as exc:
        sys.exit(f"Lecture de Sage impossible : {exc}")
    for x in anomalies[:40]:
        print(f"  {'BLOQUANT ' if x.bloquante else ''}{x.article.numero:<14} {x.article.description[:38]:38} "
              f"{x.probleme} -> {x.proposition}")
    if len(anomalies) > 40:
        print(f"  … et {len(anomalies) - 40} autre(s) : voir le fichier Excel.")


def commande_correspondance(args: argparse.Namespace) -> None:
    from decimal import Decimal

    from .api_clover import SANS_CATEGORIE, ClientClover
    from .articles_sage import lire_articles_sage
    from .correspondance import (AUCUN, AUTOMATIQUE, A_CONFIRMER, CONFIRME, ArticleClover,
                                 ecrire_correspondance, lire_confirmes, rapprocher)

    config_toml = tomllib.loads(Path(args.config).read_text(encoding="utf-8")) if Path(args.config).exists() else {}
    sage = None
    if not args.sage and config_toml.get("sage_base", {}).get("utilisateur"):
        try:
            from .articles_sage import ecrire_copie

            sage = lire_sage(config_toml, lambda base: base.articles())
            ecrire_copie(sage, RACINE / "donnees" / COPIE_ARTICLES_SAGE)
            print(f"Articles lus directement dans Sage : {len(sage)}")
            verifier_comptes_sage(config_toml, fichier_sortie=rapport_comptes(config_toml.get("sage_base", {}).get("fichier_entreprise")))
        except Exception as exc:  # Sage fermé, mot de passe refusé… : on se rabat sur le fichier exporté
            print(f"Lecture directe de Sage impossible ({exc}) : utilisation du fichier exporté.")
    if sage is None:
        chemin_sage = Path(args.sage or config_toml.get("sage", {}).get("articles", "donnees/Article.txt"))
        if not chemin_sage.is_absolute():
            chemin_sage = RACINE / chemin_sage
        if not chemin_sage.is_file():
            sys.exit("Liste des articles Sage introuvable : indiquez-la avec --sage ou dans config.toml ([sage] articles).")
        sage = lire_articles_sage(chemin_sage)
        # Un export abîmé remplace les lettres accentuées par des espaces (« Coupé » -> « Coup ») :
        # on vérifie la présence de vraies lettres accentuées dans le fichier brut.
        accents = sum(chemin_sage.read_text(encoding="cp1252").count(c) for c in "éèàçêô")
        if len(sage) > 100 and accents == 0:
            sys.exit(f"La liste Sage {chemin_sage} ne contient aucun accent : l'export semble abîmé "
                     "(« Coupé » devient « Coup »). Réexportez la liste depuis Sage, puis relancez la commande.")

    client = ClientClover(charger_config(Path(args.config)))
    print("Lecture du catalogue Clover…")
    catalogue = client.articles()
    fin = date.today() - timedelta(days=1)
    depuis = date.fromisoformat(args.depuis) if args.depuis else fin - timedelta(days=29)
    print(f"Lecture des ventes depuis le {depuis} (articles vendus, même supprimés depuis du catalogue)…")
    ventes: dict[str, Decimal] = {}
    ventes_par_id: dict[str, Decimal] = {}
    noms_vendus: dict[str, str] = {}
    for ligne in client.rapport(depuis, fin).lignes:
        ventes[ligne.nom] = ventes.get(ligne.nom, Decimal("0")) + ligne.ventes_nettes
        if ligne.id_article:
            ventes_par_id[ligne.id_article] = ventes_par_id.get(ligne.id_article, Decimal("0")) + ligne.ventes_nettes
            noms_vendus.setdefault(ligne.id_article, ligne.nom)
    # Articles vendus récemment mais supprimés du catalogue Clover depuis : il faut aussi les relier.
    ids_catalogue = {a["id"] for a in catalogue}
    from .correspondance import lire_a_relier
    a_confirmer = {}
    for id_article, nom_article, numero_sage in lire_a_relier(A_RELIER):  # articles trouvés en préparant des factures
        noms_vendus.setdefault(id_article, nom_article)
        if numero_sage:
            a_confirmer[id_article] = (nom_article, numero_sage)
    for id_article in sorted(set(noms_vendus) - ids_catalogue):
        a = client.article(id_article) or {"id": id_article}
        a["name"] = a.get("name") or noms_vendus[id_article]
        catalogue.append(a)
        print(f"  Article vendu mais absent du catalogue : {a['name']}")

    clover = [
        ArticleClover(
            id=a["id"], nom=(a.get("name") or "").strip(),
            categorie=(((a.get("categories") or {}).get("elements") or [{}])[0]).get("name", SANS_CATEGORIE),
            code=(a.get("code") or "").strip(),
            ventes_30j=ventes_par_id.get(a["id"], ventes.get((a.get("name") or "").strip(), Decimal("0"))),
        )
        for a in catalogue
    ]
    chemin = Path(args.fichier)
    confirmes = lire_confirmes(chemin)
    if chemin.exists():  # copie de sécurité du fichier validé avant de le régénérer
        import shutil
        from datetime import datetime
        copie = chemin.parent / "sauvegardes" / f"{chemin.stem}_{datetime.now():%Y%m%d-%H%M%S}{chemin.suffix}"
        copie.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(chemin, copie)
        print(f"Copie de sécurité : {copie}")
    from .correspondance import Correspondance as _C
    for id_article, (nom_article, numero_sage) in a_confirmer.items():
        confirmes[id_article] = _C(CONFIRME, nom_article, "", Decimal("0"), numero_sage, "", "",
                                   "Manuel (articles_a_relier)", 100, "", "", id_article)
    resultat = rapprocher(clover, sage, confirmes)
    try:
        ecrire_correspondance(resultat, chemin)
    except PermissionError:
        sys.exit(f"Impossible d'enregistrer {chemin.name} : fermez-le dans Excel, puis relancez la commande.")

    compte = {s: sum(1 for c in resultat if c.statut == s) for s in (CONFIRME, AUTOMATIQUE, A_CONFIRMER, AUCUN)}
    vendus = [c for c in resultat if c.ventes_30j]
    print(f"\nArticles Sage lus : {len(sage)} — articles Clover : {len(clover)}")
    for statut, nombre in compte.items():
        print(f"  {statut:<12}: {nombre:>5}   (dont vendus ces 30 jours : "
              f"{sum(1 for c in vendus if c.statut == statut)})")
    print(f"\nFichier à vérifier : {chemin}")


A_RELIER = RACINE / "donnees" / "articles_a_relier.csv"


def memoriser_a_relier(articles: list[tuple[str, str]]) -> None:
    """Articles vendus sans numéro Sage : la prochaine « correspondance » les ajoutera (même supprimés de Clover)."""
    import csv

    from .correspondance import lire_a_relier

    deja = {r[0] for r in lire_a_relier(A_RELIER)}
    nouveaux = [(i, n) for i, n in dict(articles).items() if i and i not in deja]
    if nouveaux:
        A_RELIER.parent.mkdir(parents=True, exist_ok=True)
        with A_RELIER.open("a", encoding="utf-8", newline="") as f:
            csv.writer(f, delimiter=";").writerows(nouveaux)


def config_import(config_toml: dict):
    from .import_sage import ConfigImport

    c = config_toml.get("sage_import", {})
    return ConfigImport(version=str(c.get("version", "12001")), pays=str(c.get("pays", "1")),
                        position_numero=int(c.get("position_numero", 3)))


def produire_factures(args: argparse.Namespace):
    """Prépare la ou les factures de la période (une par jour avec --par-jour) et écrit le fichier .IMP
    (toutes les factures dans un seul fichier : un seul import, une seule confirmation dans Sage) et le détail Excel.

    Retourne [(import, problèmes)], le chemin du .IMP et celui du détail Excel.
    """
    import openpyxl
    from openpyxl.styles import Font

    from .facture import preparer_facture
    from .import_sage import ecrire_imp, lire_correspondance, nom_fichier, preparer_import

    type_periode = TypePeriode(args.periode)
    debut, fin = bornes(type_periode, date.fromisoformat(args.date))
    par_jour = getattr(args, "par_jour", False) and debut != fin
    # --par-jour : une facture par jour du lundi au jeudi, une seule du vendredi au dimanche
    periodes = blocs_facturation(debut, fin) if par_jour else [(debut, fin, type_periode)]

    chemin_corresp = Path(args.correspondance)
    if not chemin_corresp.exists():
        sys.exit(f"Fichier de correspondance introuvable : {chemin_corresp} (lancez « correspondance »).")
    table = lire_correspondance(chemin_corresp)
    client = None
    if not getattr(args, "export", None):
        from .api_clover import ClientClover
        client = ClientClover(charger_config(Path(args.config)))

    resultats = []
    for d, f, type_facture in periodes:
        if client is None:
            from .export_clover import lire_export
            rapport = lire_export(Path(args.export), d, f)
        else:
            rapport = client.rapport(d, f)
        if not rapport.lignes:
            print(f"{d} : aucune vente, pas de facture.")
            continue
        facture = preparer_facture(rapport, type_facture)
        imp = preparer_import(facture, table)
        problemes = [f"Non relié à Sage (non importé) : {n}" for n in imp.non_reliees] + facture.anomalies
        if imp.sous_total != rapport.total_net:
            problemes.insert(0, f"Écart de {rapport.total_net - imp.sous_total:,.2f} $ avec les ventes nettes Clover")
        resultats.append((imp, problemes, rapport))
        memoriser_a_relier(imp.ids_non_relies)
        print(f"{facture.numero:<22} {len(imp.lignes):>4} lignes  ventes nettes {rapport.total_net:>10,.2f} $  "
              f"sous-total {imp.sous_total:>10,.2f} $  taxes {imp.total_taxes:>7,.2f} $  total {imp.total:>10,.2f} $"
              + (f"  ! {len(problemes)} point(s) à vérifier" if problemes else ""))
        for texte in problemes:
            print(f"    ! {texte}")
        for texte in facture.avertissements:
            print(f"    (info) {texte}")
    if not resultats:
        sys.exit("Aucune vente sur la période : aucune facture à préparer.")

    nom = nom_fichier(resultats[0][0].facture.numero) if len(resultats) == 1 else f"Ventes_{debut}_au_{fin}.IMP"
    chemin_imp = ecrire_imp([r[0] for r in resultats], RACINE / "sortie" / "factures" / nom,
                            config_import(lire_toml(args.config)))

    classeur = openpyxl.Workbook()
    resume = classeur.active
    resume.title = "Résumé"
    resume.append(["Facture", "Date", "Lignes", "Ventes nettes Clover", "Sous-total", "Taxes", "Total",
                   "Points à vérifier"])
    for imp, problemes, rapport in resultats:
        resume.append([imp.facture.numero, imp.facture.date_facture, len(imp.lignes), float(rapport.total_net),
                       float(imp.sous_total), float(imp.total_taxes), float(imp.total), " | ".join(problemes)])
    resume.append([])
    resume.append(["Total", "", sum(len(r[0].lignes) for r in resultats),
                   float(sum(r[2].total_net for r in resultats)), float(sum(r[0].sous_total for r in resultats)),
                   float(sum(r[0].total_taxes for r in resultats)), float(sum(r[0].total for r in resultats))])
    for ligne in (resume[1], resume[resume.max_row]):
        for cellule in ligne:
            cellule.font = Font(bold=True)
    for colonne, largeur in zip("ABCDEFGH", (22, 12, 8, 18, 14, 10, 14, 80)):
        resume.column_dimensions[colonne].width = largeur
    for imp, _, _ in resultats:
        feuille = classeur.create_sheet(imp.facture.date_facture.isoformat() if par_jour else "Facture")
        feuille.append(["N° article Sage", "Articles Clover", "Quantité", "Prix", "Montant", "Taxe", "TPS/TVH", "TVQ"])
        for cellule in feuille[1]:
            cellule.font = Font(bold=True)
        for l in imp.lignes:
            t = {n: m for n, _, m in l.taxes()}
            feuille.append([l.numero_sage, " | ".join(l.noms_clover), float(l.quantite), float(l.prix),
                            float(l.montant), "GQ" if l.taxable else "", float(t.get("TPS/TVH", 0)),
                            float(t.get("TVQ", 0))])
        feuille.column_dimensions["B"].width = 60
    chemin_xlsx = chemin_imp.with_suffix(".xlsx")
    classeur.save(chemin_xlsx)
    print(f"\nFichier d'import : {chemin_imp}\nDétail (Excel)   : {chemin_xlsx}")
    return [(r[0], r[1]) for r in resultats], chemin_imp, chemin_xlsx


def lire_toml(chemin) -> dict:
    chemin = Path(chemin)
    return tomllib.loads(chemin.read_text(encoding="utf-8")) if chemin.exists() else {}


def commande_facture(args: argparse.Namespace) -> None:
    produire_factures(args)


def executable_import() -> Path | None:
    """Sage_SA_import.exe de la version de Sage 50 la plus récente installée."""
    candidats = sorted(Path("C:/Program Files (x86)").glob("Sage 50*/Sage_SA_import.exe"), reverse=True)
    return candidats[0] if candidats else None


def commande_importer(args: argparse.Namespace) -> None:
    import csv
    import subprocess
    from datetime import datetime

    from .import_sage import ecrire_imp

    config_toml = lire_toml(args.config)
    c = config_toml.get("sage_import", {})
    fichier_sai = Path(c.get("fichier_entreprise", ""))
    if not c.get("fichier_entreprise") or fichier_sai.suffix.upper() != ".SAI" or not fichier_sai.is_file():
        sys.exit("Indiquez le fichier d'entreprise Sage (.SAI) dans config.toml : [sage_import] fichier_entreprise.")
    exe = Path(c["executable"]) if c.get("executable") else executable_import()
    if not exe or not exe.is_file():
        sys.exit("Sage_SA_import.exe introuvable : indiquez-le dans config.toml ([sage_import] executable).")

    factures, chemin_imp, _ = produire_factures(args)
    client = factures[0][0].facture.client
    journal = RACINE / "donnees" / "journal_imports.csv"

    def noter(imp, code_retour, statut):
        nouveau = not journal.exists()
        with journal.open("a", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            if nouveau:
                w.writerow(["Date et heure", "Facture", "Fichier d'entreprise", "Lignes", "Sous-total", "Total",
                            "Code de retour", "Résultat"])
            w.writerow([f"{datetime.now():%Y-%m-%d %H:%M:%S}", imp.facture.numero, fichier_sai.name,
                        len(imp.lignes), f"{imp.sous_total:.2f}", f"{imp.total:.2f}", code_retour, statut])

    # Contrôles dans le fichier d'entreprise visé
    try:
        with ouvrir_base_sage(config_toml, fichier_sai) as base:
            mes_articles = base.client_utilise_mes_articles(client)
            numeros = sorted({l.numero_sage for imp, _ in factures for l in imp.lignes})
            sans_compte = base.articles_sans_compte_revenus(numeros)
            invisibles = {n: d for n, d in base.numeros_avec_caracteres_invisibles().items() if n in set(numeros)}
            existantes = {imp.facture.numero for imp, _ in factures if base.facture_existe(imp.facture.numero, client)}
    except Exception as exc:
        if not args.forcer:
            sys.exit(f"\nImport annulé : vérifications impossibles dans Sage ({exc}). "
                     "Ouvrez le fichier d'entreprise dans Sage ou relancez avec --forcer.")
        mes_articles, sans_compte, existantes, invisibles = True, [], set(), {}
    if mes_articles is None:
        sys.exit(f"\nImport annulé : le client « {client} » n'existe pas dans {fichier_sai.name}.")
    if not mes_articles:
        sys.exit(f"\nImport annulé : dans {fichier_sai.name}, la fiche du client « {client} » (onglet "
                 "Importer/Exporter) doit indiquer que le client utilise vos numéros d'articles. "
                 "Sinon Sage ne reconnaît aucun article de la facture.")
    if invisibles:
        details = "\n".join(f"   {numero}  {nom}" for numero, nom in invisibles.items())
        sys.exit(f"\nImport annulé : dans {fichier_sai.name}, le numéro de ces articles contient des caractères "
                 f"invisibles (retour à la ligne ou espace) ; Sage ne les reconnaîtrait pas. Dans la fiche de "
                 f"l'article, effacez le numéro et retapez-le :\n{details}")
    if sans_compte:
        details = "\n".join(f"   {numero}  {nom}" for numero, nom in sans_compte)
        sys.exit(f"\nImport annulé : ces articles n'ont pas de compte de revenus dans {fichier_sai.name} "
                 f"(fiche de l'article dans Sage) :\n{details}")
    try:  # comptes incohérents : avertissement seulement (Sage importe quand même)
        for x in verifier_comptes_sage(config_toml, fichier_sai, set(numeros)):
            if not x.bloquante:
                print(f"  ! {x.article.numero} {x.article.description} : {x.probleme}")
    except Exception:
        pass

    a_importer = []
    for imp, problemes in factures:
        if imp.facture.numero in existantes:
            print(f"  {imp.facture.numero} : existe déjà dans Sage, ignorée.")
            noter(imp, "", "déjà présente")
        elif problemes and not args.forcer:
            print(f"  {imp.facture.numero} : points à vérifier, non importée (relancer avec --forcer pour l'importer).")
            noter(imp, "", "non importée : points à vérifier")
        else:
            a_importer.append(imp)
    if not a_importer:
        sys.exit("\nAucune facture à importer.")
    if len(a_importer) != len(factures):
        chemin_imp = ecrire_imp(a_importer, chemin_imp, config_import(config_toml))

    print(f"\nImport de {len(a_importer)} facture(s) dans {fichier_sai.name} avec {exe.name}…")
    try:
        code_retour = subprocess.run([str(exe), str(chemin_imp), str(fichier_sai)], timeout=args.delai).returncode
    except subprocess.TimeoutExpired:  # Sage attend une réponse dans une fenêtre : le processus est arrêté
        code_retour = "délai dépassé"
        print(f"  Import interrompu après {args.delai} s (Sage attendait probablement une réponse dans une fenêtre).")

    echecs = 0
    try:
        with ouvrir_base_sage(config_toml, fichier_sai) as base:
            presentes = {imp.facture.numero: base.facture_existe(imp.facture.numero, client) for imp in a_importer}
    except Exception as exc:
        print(f"  (vérification dans Sage impossible : {exc})")
        presentes = {imp.facture.numero: None for imp in a_importer}
    for imp in a_importer:
        present = presentes[imp.facture.numero]
        statut = "importée" if present else ("non vérifiable" if present is None else "ÉCHEC")
        echecs += present is False
        noter(imp, code_retour, statut)
        print(f"  {imp.facture.numero:<22} {imp.total:>10,.2f} $  {statut}")
    print(f"Journal : {journal}")
    if echecs:
        sys.exit(f"ÉCHEC : {echecs} facture(s) introuvable(s) dans Sage après l'import (code {code_retour}).")


def commande_auto(args: argparse.Namespace) -> None:
    from .automatique import lancer

    cr = lancer(RACINE, lire_toml(args.config), date.fromisoformat(args.jusqua) if args.jusqua else None,
                simulation=args.simulation)
    for p in cr.problemes:
        print(f"PROBLÈME : {p}")
    if cr.erreur:
        print(f"ERREUR : {cr.erreur}")
    if not cr.succes:
        sys.exit(1)


def commande_planifier(args: argparse.Namespace) -> None:
    from . import planification

    try:
        if args.retirer:
            print(planification.retirer())
        elif args.etat:
            print(planification.etat())
        else:
            print(planification.installer(RACINE, args.heure))
    except RuntimeError as exc:
        sys.exit(f"Planification impossible : {exc}")


def commande_langue(args: argparse.Namespace) -> None:
    """Enregistre la langue de l'application (appelée par l'installateur avec la langue choisie)."""
    from .interface import ecrire_config, lire_config
    from .langues import LANGUES

    code = args.code.strip().lower()[:2]
    if code not in LANGUES:
        sys.exit(f"Langue inconnue : {args.code} (choix : {', '.join(LANGUES)})")
    chemin = Path(args.config)
    config = lire_config(chemin)
    config.setdefault("general", {})["langue"] = code
    ecrire_config(chemin, config)
    print(f"Langue de l'application : {LANGUES[code]}")


def commande_interface(args: argparse.Namespace) -> None:
    from .interface import lancer_interface

    lancer_interface(RACINE)


def main(argv: list[str] | None = None) -> None:
    if argv is None and len(sys.argv) == 1:
        argv = ["interface"]  # double-clic sur le programme : ouvrir l'interface
    for flux in (sys.stdout, sys.stderr):
        if flux is not None and hasattr(flux, "reconfigure"):
            flux.reconfigure(encoding="utf-8", errors="replace")
    parseur = argparse.ArgumentParser(prog="mvogmbi_sync", description="Synchronisation Clover → Sage 50")
    sous = parseur.add_subparsers(dest="commande", required=True)

    p = sous.add_parser("rapport", help="Cumul des ventes par article pour un jour, une semaine ou un mois")
    p.add_argument("--periode", choices=[t.value for t in TypePeriode], default="jour")
    p.add_argument("--date", default=date.today().isoformat(),
                   help="Un jour compris dans la période (AAAA-MM-JJ). Par défaut : aujourd'hui.")
    p.add_argument("--par-jour", action="store_true", help="Un cumul séparé pour chaque jour de la période")
    p.add_argument("--export", help="Lire un export Excel Clover au lieu de l'API")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_rapport)

    p = sous.add_parser("verifier", help="Vérifier la connexion à Clover (config.toml)")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_verifier)

    p = sous.add_parser("comparer", help="Comparer l'API Clover avec un export journalier de la même date")
    p.add_argument("--date", required=True, help="Jour de l'export (AAAA-MM-JJ)")
    p.add_argument("--export", required=True, help="Export Excel Clover « Ventes d'articles » de ce jour")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_comparer)

    p = sous.add_parser("facture", help="Préparer le fichier d'import Sage (.IMP) de la facture de vente")
    p.add_argument("--periode", choices=[t.value for t in TypePeriode], default="jour")
    p.add_argument("--date", default=(date.today() - timedelta(days=1)).isoformat(),
                   help="Un jour compris dans la période (AAAA-MM-JJ). Par défaut : hier.")
    p.add_argument("--export", help="Lire un export Clover au lieu de l'API (sans information de taxe)")
    p.add_argument("--correspondance", default=str(RACINE / "donnees" / "correspondance_articles.xlsx"))
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.add_argument("--par-jour", action="store_true",
                   help="Une facture par jour du lundi au jeudi et une pour vendredi-dimanche "
                        "(ex. tout un mois : --periode mois --par-jour)")
    p.set_defaults(fonction=commande_facture)

    p = sous.add_parser("importer", help="Préparer la facture et l'importer automatiquement dans Sage")
    p.add_argument("--periode", choices=[t.value for t in TypePeriode], default="jour")
    p.add_argument("--date", default=(date.today() - timedelta(days=1)).isoformat(),
                   help="Un jour compris dans la période (AAAA-MM-JJ). Par défaut : hier.")
    p.add_argument("--correspondance", default=str(RACINE / "donnees" / "correspondance_articles.xlsx"))
    p.add_argument("--forcer", action="store_true", help="Importer malgré les points à vérifier")
    p.add_argument("--delai", type=int, default=900, help="Durée maximale de l'import, en secondes")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.add_argument("--par-jour", action="store_true",
                   help="Une facture par jour du lundi au jeudi et une pour vendredi-dimanche "
                        "(ex. tout un mois : --periode mois --par-jour)")
    p.set_defaults(fonction=commande_importer)

    p = sous.add_parser("langue", help="Choisir la langue de l'application : fr ou en")
    p.add_argument("code", help="fr (français) ou en (English)")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_langue)

    p = sous.add_parser("interface", help="Ouvrir l'application de configuration et de suivi")
    p.set_defaults(fonction=commande_interface)

    p = sous.add_parser("auto", help="Traitement automatique : de la dernière facture dans Sage jusqu'à hier")
    p.add_argument("--jusqua", help="Dernier jour à traiter (AAAA-MM-JJ) ; par défaut hier")
    p.add_argument("--simulation", action="store_true", help="Tout préparer et vérifier, sans importer")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_auto)

    p = sous.add_parser("planifier", help="Installer, consulter ou retirer la tâche planifiée Windows")
    p.add_argument("--heure", default="01:00")
    p.add_argument("--retirer", action="store_true")
    p.add_argument("--etat", action="store_true")
    p.set_defaults(fonction=commande_planifier)

    p = sous.add_parser("verifier-comptes", help="Vérifier les comptes des articles Sage (15xx -> 40xx / 50xx)")
    p.add_argument("--fichier", help="Fichier d'entreprise Sage (.SAI) ; par défaut celui de [sage_base]")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_verifier_comptes)

    p = sous.add_parser("encaissements", help="Dépôts Clover à encaisser sur les factures et leur état dans Sage")
    p.add_argument("--du", default=(date.today() - timedelta(days=31)).isoformat(), help="Premier jour de vente")
    p.add_argument("--au", default=(date.today() - timedelta(days=1)).isoformat(), help="Dernier jour de vente")
    p.add_argument("--export", help="Fichier Excel à créer (liste des encaissements)")
    p.add_argument("--appliquer", action="store_true",
                   help="Saisir les dépôts « À saisir » dans Sage (Sage ouvert, fenêtre Encaissements affichée)")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_encaissements)

    p = sous.add_parser("encaisser-bureau", help="(interne) Encaisser les dépôts dans Sage sur le bureau caché")
    p.add_argument("--resultat", required=True, help="Fichier JSON du bilan")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_encaisser_bureau)

    p = sous.add_parser("verifier-sage", help="Vérifier la lecture directe des données de Sage 50")
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.set_defaults(fonction=commande_verifier_sage)

    p = sous.add_parser("correspondance", help="Créer ou mettre à jour la correspondance articles Clover -> Sage")
    p.add_argument("--sage", help="Liste des articles exportée de Sage (ex. E:\\Article.txt)")
    p.add_argument("--fichier", default=str(RACINE / "donnees" / "correspondance_articles.xlsx"))
    p.add_argument("--config", default=str(RACINE / "config.toml"))
    p.add_argument("--depuis", help="Prendre en compte les articles vendus depuis cette date (AAAA-MM-JJ) ; "
                   "par défaut les 30 derniers jours")
    p.set_defaults(fonction=commande_correspondance)

    args = parseur.parse_args(argv)
    args.fonction(args)


if __name__ == "__main__":
    main()
