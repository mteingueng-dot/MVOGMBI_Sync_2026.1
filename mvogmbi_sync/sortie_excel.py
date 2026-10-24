"""Écriture d'un rapport de ventes cumulé dans un fichier Excel de vérification."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from .models import RapportVentes

ENTETES = [
    "Catégorie", "Article", "Qté vendue", "Qté remboursée", "Qté nette",
    "Ventes brutes", "Rabais", "Remboursements", "Ventes nettes", "Prix unitaire net",
]
FORMAT_ARGENT = '#,##0.00 "$"'


def ecrire_rapports(rapports: list[RapportVentes], chemin: Path) -> Path:
    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)
    for rapport in rapports:
        nom = rapport.debut.isoformat() if rapport.debut == rapport.fin else f"{rapport.debut}_{rapport.fin}"
        feuille = classeur.create_sheet(nom[:31])
        feuille.append([f"Ventes du {rapport.debut} au {rapport.fin} — {rapport.source}"])
        feuille.append(ENTETES)
        for cellule in feuille[2]:
            cellule.font = Font(bold=True)
        for l in rapport.lignes:
            feuille.append([
                l.categorie, l.nom, l.quantite_vendue, l.quantite_remboursee, l.quantite_nette,
                l.ventes_brutes, l.rabais, l.remboursements, l.ventes_nettes, l.prix_unitaire_net,
            ])
        feuille.append([])
        feuille.append(["TOTAL", "", "", "", "", rapport.total_brut, rapport.total_rabais,
                        rapport.total_remboursements, rapport.total_net])
        for cellule in feuille[feuille.max_row]:
            cellule.font = Font(bold=True)
        for colonne in "FGHIJ":
            for cellule in feuille[colonne][2:]:
                cellule.number_format = FORMAT_ARGENT
        feuille.column_dimensions["A"].width = 24
        feuille.column_dimensions["B"].width = 48
        for colonne in "CDEFGHIJ":
            feuille.column_dimensions[colonne].width = 15
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin
